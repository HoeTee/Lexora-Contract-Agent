import asyncio
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from docx import Document

from agents.orchestrator import OrchestratorAgent
from agents.schemas import SubAgentOutput
from routing import load_mandatory_tasks
from tools.document.reporting.docx_report import deduplicate_review_results_by_range
from workflow.workflow import ContractReviewWorkflow


class MandatoryRoutingTests(unittest.TestCase):
    def test_yaml_profile_defines_all_18_fixed_routes(self) -> None:
        tasks = load_mandatory_tasks()

        self.assertEqual(len(tasks), 18)
        self.assertEqual([task["id"] for task in tasks], [f"H{i:02d}" for i in range(1, 19)])
        self.assertTrue(all(task["review_source"] == "mandatory" for task in tasks))
        self.assertTrue(all(not task["fallback_llm_route"] for task in tasks))
        self.assertEqual(tasks[11]["dedupe_domain"], "language_quality")
        self.assertEqual(tasks[14]["dedupe_domain"], "language_quality")
        self.assertEqual(tasks[10]["evidence_region"], "attachments")
        methods = {step["method"] for task in tasks for step in task["route"]["steps"]}
        self.assertEqual(methods, {"title", "region", "rule", "scan", "hybrid"})

    def test_h01_treats_each_explicit_empty_field_as_a_local_defect(self) -> None:
        task = load_mandatory_tasks()[0]

        self.assertIn("相邻位置存在另一个同名字段或实际值不能推翻", task["check_points"][0])

    def test_h07_scopes_date_and_term_checks_without_attachments(self) -> None:
        task = load_mandatory_tasks()[6]

        self.assertEqual(
            task["route"]["steps"],
            [
                {
                    "method": "region",
                    "terms": [],
                    "regions": ["frontmatter"],
                    "all_titles": False,
                    "anchor_terms": [],
                    "coverage": "complete",
                    "candidate_instruction": "",
                },
                {
                    "method": "region",
                    "terms": ["合同期限", "履行期限", "服务期限", "交付期限"],
                    "regions": ["body"],
                    "all_titles": False,
                    "anchor_terms": [],
                    "coverage": "relevant",
                    "candidate_instruction": "",
                },
            ],
        )
        joined_check_points = "\n".join(task["check_points"])
        self.assertIn("必须分别生成两个 issue", joined_check_points)
        self.assertIn("不得借用合同名称、主体名称", joined_check_points)

    def test_h06_only_section_level_problem_uses_payment_title(self) -> None:
        task = load_mandatory_tasks()[5]

        self.assertEqual(task["placement"]["target"], "inline")
        joined_check_points = "\n".join(task["check_points"])
        self.assertIn("同一个标题 edit 最多承载一个章节整体问题", joined_check_points)
        self.assertIn("必须锚定直接包含该问题的具体原文", joined_check_points)

    def test_h18_compares_body_references_with_titles_without_attachment_bodies(self) -> None:
        task = load_mandatory_tasks()[17]

        self.assertEqual(
            task["route"]["steps"],
            [
                {
                    "method": "scan",
                    "terms": [],
                    "regions": ["body"],
                    "all_titles": False,
                    "anchor_terms": [],
                    "coverage": "relevant",
                    "candidate_instruction": "提取本区域中每一处明确引用附件编号、附件名称或附件简称的原文；只要构成附件引用就返回，不要求本批次内已经出现实际附件标题。",
                },
                {
                    "method": "title",
                    "terms": [],
                    "regions": ["attachments"],
                    "all_titles": True,
                    "anchor_terms": [],
                    "coverage": "relevant",
                    "candidate_instruction": "",
                },
            ],
        )
        self.assertTrue(task["comparison_audit"])
        self.assertEqual(task["evidence_region"], "attachments")

    def test_h18_comparison_audit_requires_complete_paired_decisions(self) -> None:
        task = load_mandatory_tasks()[17]
        context = """source_region: body
xml_anchor_type: paragraph
xml_anchor_id: path:body/p70
内容：
详见附件一《产品清单》

---

source_region: attachments
xml_anchor_type: paragraph
xml_anchor_id: path:body/p219
内容：
附件一：服务清单"""
        missing = OrchestratorAgent._policy_output_errors(
            {"issues": [], "comparison_decisions": []}, task, context
        )
        self.assertTrue(any("缺少正文引用锚点" in error for error in missing))

        valid = SubAgentOutput.model_validate({
            "status": "issues_found",
            "issues": [{
                "issue_id": "H18.1",
                "risk_level": "low",
                "issue_comment": "正文引用的附件名称与实际附件标题不一致。",
                "anchors": [
                    {"role": "edit", "defect_type": "inconsistency", "xml_anchor_type": "paragraph", "xml_anchor_id": "path:body/p70", "quoted_text": "附件一《产品清单》", "comment_text": "请将正文附件名称修改为“服务清单”。"},
                    {"role": "evidence", "defect_type": "inconsistency", "xml_anchor_type": "paragraph", "xml_anchor_id": "path:body/p219", "quoted_text": "附件一：服务清单", "comment_text": ""},
                ],
                "reasoning": "正文名称与实际附件标题不同。",
                "criterion": task["criterion"],
                "check_point": task["check_points"][0],
            }],
            "coverage_decisions": [],
            "comparison_decisions": [{
                "edit_anchor_type": "paragraph",
                "edit_anchor_id": "path:body/p70",
                "evidence_anchor_type": "paragraph",
                "evidence_anchor_id": "path:body/p219",
                "decision": "mismatch",
                "reason": "正文为产品清单，实际附件为服务清单。",
            }],
        }).model_dump(mode="json")
        self.assertEqual(OrchestratorAgent._policy_output_errors(valid, task, context), [])

    def test_h18_allows_one_paragraph_to_reference_multiple_attachments(self) -> None:
        base = {
            "edit_anchor_type": "paragraph",
            "edit_anchor_id": "path:body/p1",
            "evidence_anchor_type": "paragraph",
            "decision": "mismatch",
        }
        output = SubAgentOutput.model_validate({
            "status": "issues_found",
            "issues": [{
                "issue_id": "H18.1",
                "risk_level": "low",
                "issue_comment": "同一段正文引用的附件名称与实际标题不一致。",
                "anchors": [{
                    "role": "edit",
                    "defect_type": "inconsistency",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "附件2和附件3",
                    "comment_text": "请统一附件名称。",
                }],
                "reasoning": "同一段引用两个附件。",
                "criterion": "附件名称",
                "check_point": "逐一比较",
            }],
            "comparison_decisions": [
                {**base, "evidence_anchor_id": "path:body/p2", "reason": "附件2不一致"},
                {**base, "evidence_anchor_id": "path:body/p3", "reason": "附件3不一致"},
            ],
        })

        self.assertEqual(len(output.comparison_decisions), 2)

    def test_h18_degraded_cleanup_preserves_comparison_issues(self) -> None:
        parsed = {"status": "issues_found", "issues": [{"issue_id": "H18.1"}]}

        removed = OrchestratorAgent._drop_incomplete_policy_issues(
            parsed,
            {"comparison_audit": True, "evidence_region": "attachments"},
            "",
        )

        self.assertEqual(removed, 0)
        self.assertEqual([item["issue_id"] for item in parsed["issues"]], ["H18.1"])

    def test_mandatory_only_phase_skips_planner_model_call(self) -> None:
        workflow = ContractReviewWorkflow.__new__(ContractReviewWorkflow)
        workflow.state_store = None
        workflow.settings = None
        workflow.logger = SimpleNamespace(log=lambda **_kwargs: None)

        with patch("workflow.workflow.PlannerAgent", side_effect=AssertionError("Planner must be skipped")):
            tasks = asyncio.run(workflow._phase_plan("", None))

        self.assertEqual(len(tasks), 18)
        self.assertEqual(workflow._planner_tokens, 0)

    def test_uploaded_criteria_are_appended_as_planned_tasks(self) -> None:
        workflow = ContractReviewWorkflow.__new__(ContractReviewWorkflow)
        workflow.state_store = None
        workflow.settings = None
        workflow.logger = SimpleNamespace(log=lambda **_kwargs: None)
        planner = SimpleNamespace(
            extract_strategies=AsyncMock(return_value={
                "U01": {
                    "route": {"steps": [{"method": "scan"}], "fallback": False},
                    "placement": {"target": "inline", "note": ""},
                }
            }),
            token_usage={"total_tokens": 12},
        )
        skeletons = [{"number": 1, "section": "用户标准", "criterion": "检查自定义事项"}]

        with (
            patch("workflow.workflow.split_criteria_from_docx", return_value=skeletons),
            patch("workflow.workflow.PlannerAgent", return_value=planner),
        ):
            tasks = asyncio.run(workflow._phase_plan("custom", "uploaded-criteria.docx"))

        self.assertEqual(len(tasks), 19)
        self.assertEqual(tasks[-1]["id"], "U01")
        self.assertEqual(tasks[-1]["review_source"], "uploaded")
        planner.extract_strategies.assert_awaited_once()
        self.assertEqual(workflow._planner_tokens, 12)

    def test_same_resolved_range_prefers_mandatory_over_uploaded(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "contract.docx"
            document = Document()
            document.add_paragraph("乙方：#乙方名称#")
            document.save(path)
            mandatory = _result("H01", "mandatory", "请填写乙方名称。")
            uploaded = _result("U01", "uploaded", "请核实乙方名称。")

            results, stats = deduplicate_review_results_by_range(
                str(path),
                [uploaded, mandatory],
            )

        self.assertEqual(results[0]["issues"], [])
        self.assertEqual(len(results[1]["issues"]), 1)
        self.assertEqual(stats["removed_duplicate_issues"], 1)

    def test_same_source_exact_range_keeps_higher_priority_finding(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "contract.docx"
            document = Document()
            document.add_paragraph("乙方：#乙方名称#")
            document.save(path)
            first = _result("H01", "mandatory", "请填写乙方名称。")
            second = _result("H10", "mandatory", "乙方名称未填写。")
            first["dedupe_priority"] = 100
            second["dedupe_priority"] = 200

            results, stats = deduplicate_review_results_by_range(str(path), [first, second])

        self.assertEqual([len(result["issues"]) for result in results], [0, 1])
        self.assertEqual(stats["removed_duplicate_issues"], 1)

    def test_uploaded_exact_range_never_overrides_mandatory(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "contract.docx"
            document = Document()
            document.add_paragraph("乙方：#乙方名称#")
            document.save(path)
            mandatory = _result("H01", "mandatory", "请填写乙方名称。")
            uploaded = _result("U01", "uploaded", "请核实乙方名称。")
            uploaded["dedupe_priority"] = 999

            results, _ = deduplicate_review_results_by_range(
                str(path),
                [uploaded, mandatory],
            )

        self.assertEqual(results[0]["issues"], [])
        self.assertEqual(len(results[1]["issues"]), 1)

    def test_contained_ranges_in_same_domain_keep_higher_priority(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "contract.docx"
            document = Document()
            document.add_paragraph("不得贿赂贵方项目和管理人员。")
            document.save(path)
            typo = _result("H12", "mandatory", "项目不能作为贿赂对象。")
            grammar = _result("H15", "mandatory", "建议改为项目人员和管理人员。")
            typo["issues"][0]["anchors"][0]["quoted_text"] = "贿赂贵方项目和管理人员。"
            grammar["issues"][0]["anchors"][0]["quoted_text"] = "不得贿赂贵方项目和管理人员。"
            typo.update(dedupe_priority=100, dedupe_domain="language_quality")
            grammar.update(dedupe_priority=200, dedupe_domain="language_quality")

            results, stats = deduplicate_review_results_by_range(str(path), [typo, grammar])

        self.assertEqual([len(result["issues"]) for result in results], [0, 1])
        self.assertEqual(stats["removed_overlap_anchors"], 1)

    def test_contained_ranges_in_different_domains_are_preserved(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "contract.docx"
            document = Document()
            document.add_paragraph("合同金额人民币捌万元整。")
            document.save(path)
            first = _result("H12", "mandatory", "文字问题。")
            second = _result("H14", "mandatory", "金额问题。")
            first["issues"][0]["anchors"][0]["quoted_text"] = "人民币捌万元整"
            second["issues"][0]["anchors"][0]["quoted_text"] = "合同金额人民币捌万元整。"
            first["dedupe_domain"] = "language_quality"
            second["dedupe_domain"] = "amount"

            results, stats = deduplicate_review_results_by_range(str(path), [first, second])

        self.assertEqual([len(result["issues"]) for result in results], [1, 1])
        self.assertEqual(stats["removed_overlap_anchors"], 0)


def _result(criterion_id: str, source: str, comment: str) -> dict:
    return {
        "criterion_id": criterion_id,
        "review_source": source,
        "issues": [{
            "issue_id": f"{criterion_id}.1",
            "risk_level": "low",
            "issue_comment": comment,
            "anchors": [{
                "role": "edit",
                "defect_type": "missing_value",
                "xml_anchor_type": "paragraph",
                "xml_anchor_id": "path:body/p1",
                "quoted_text": "乙方：#乙方名称#",
                "comment_text": comment,
            }],
        }],
    }


if __name__ == "__main__":
    unittest.main()
