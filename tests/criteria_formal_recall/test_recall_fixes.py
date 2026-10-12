from __future__ import annotations

import asyncio
import inspect
import json
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from docx.oxml import OxmlElement

from agents.base_agent import estimate_chat_prompt_tokens
from agents.json_utils import AgentOutputValidationError
from agents.orchestrator import (
    OrchestratorAgent,
    _FindingLedger,
    _preserve_revision_issue_ids,
)
from agents.prompts.cn_prompts import load_prompts
from agents.schemas import ReflectorOutput, SubAgentOutput, SummaryOutput
from agents.summarizer import (
    SUMMARIZER_MAX_PROMPT_TOKENS,
    SummarizerAgent,
    apply_summary_duplicate_groups,
    conservative_message_token_upper_bound,
    normalize_summary_output,
)
from tests.criteria_formal_recall.prepare_cases import select_contract_title_group
from tools.document.docx_anchor_index import build_docx_anchor_nodes
from tools.document.reporting.docx_report import DocxReportGenerator
from tools.docxindex.llm.schemas import RouteResponse
from tools.docxindex.llm.prompts import scan_prompt
from tools.docxindex.parser import read_docx_items
from tools.docxindex.retriever import _format_anchor_context, _should_use_fallback
from tools.docxindex.retrieval.rule import _literal_terms
from workflow.workflow import _incomplete_review_notices, _prepare_inline_results


def _subagent_output(issue_id: str, comment: str, anchor_id: str) -> dict:
    return {
        "status": "issues_found",
        "applicability_reason": "",
        "issues": [{
            "issue_id": issue_id,
            "risk_level": "low",
            "issue_comment": comment,
            "anchors": [{
                "role": "edit",
                "xml_anchor_type": "paragraph",
                "xml_anchor_id": anchor_id,
                "quoted_text": "地址：【未填写】",
                "comment_text": comment,
            }],
            "reasoning": "原文字段没有填写。",
            "criterion": "检查字段是否填写。",
            "check_point": "逐字段检查。",
        }],
        "coverage_decisions": [],
    }


class AnchorEvidenceValidationTests(unittest.TestCase):
    def test_duplicate_edit_locations_are_left_for_semantic_deduplication(self) -> None:
        anchor = {
            "role": "edit",
            "xml_anchor_type": "paragraph",
            "xml_anchor_id": "path:body/p31",
            "quoted_text": "乙方：某公司 地址：【未填写】",
            "comment_text": "请填写该字段。",
        }
        issues = []
        for issue_id in ("H01.1", "H01.2"):
            issues.append({
                "issue_id": issue_id,
                "risk_level": "low",
                "issue_comment": "字段未填写。",
                "anchors": [anchor],
                "reasoning": "原文为占位内容。",
                "criterion": "基础信息齐全。",
                "check_point": "逐字段检查。",
            })

        output = SubAgentOutput.model_validate({"status": "issues_found", "issues": issues})

        self.assertEqual([issue.issue_id for issue in output.issues], ["H01.1", "H01.2"])

    def test_defect_type_is_not_part_of_anchor_schema(self) -> None:
        output = _subagent_output("H01.1", "请填写地址。", "path:body/p31")
        output["issues"][0]["anchors"][0]["defect_type"] = "missing_value"

        with self.assertRaisesRegex(ValueError, "defect_type"):
            SubAgentOutput.model_validate(output)

    def test_blank_field_examples_are_semantic_scan_instructions(self) -> None:
        prompt = scan_prompt(
            "检查合同所有字段是否填写完整",
            [{"anchor_id": "p_0001", "text": "乙方支持联络电话：XX，XX"}],
        )

        self.assertIn("结合字段名称、相邻文字、表格关系和合同语境判断", prompt)
        self.assertIn("#乙方合同号#【     】", prompt)
        self.assertIn("违约金【5%】", prompt)
        self.assertIn("条件性说明不能证明条件已经发生", prompt)
        self.assertIn("列举项不是封闭清单", prompt)
        self.assertIn("不设 Top-K", prompt)

    def test_rule_terms_are_only_planner_supplied_literals(self) -> None:
        self.assertCountEqual(_literal_terms(["联系电话", " 合同编号 "]), ["联系电话", "合同编号"])

    def test_reflector_requires_unique_per_finding_decisions(self) -> None:
        value = {
            "status": "PASS",
            "feedback": "",
            "finding_decisions": [
                {"finding_id": "F001", "decision": "confirm", "reason": "字段明确未填写"}
            ],
        }
        self.assertEqual(ReflectorOutput.model_validate(value).finding_decisions[0].decision, "confirm")
        value["finding_decisions"].append(value["finding_decisions"][0])
        with self.assertRaises(ValueError):
            ReflectorOutput.model_validate(value)

    def test_missing_fields_require_one_issue_and_one_edit_each(self) -> None:
        issue = {
            "issue_id": "C3.1",
            "risk_level": "low",
            "issue_comment": "两个字段未填写。",
            "anchors": [
                {
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "合同编号：待填写",
                    "comment_text": "请填写合同编号。",
                },
                {
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p2",
                    "quoted_text": "联系电话：待填写",
                    "comment_text": "请填写联系电话。",
                },
            ],
            "reasoning": "字段没有实际内容。",
            "criterion": "信息应填写完整。",
            "check_point": "信息应填写完整。",
        }

        with self.assertRaisesRegex(ValueError, "one separate issue per annotation location"):
            SubAgentOutput.model_validate({"status": "issues_found", "issues": [issue]})

    def test_subagent_issue_ids_must_be_unique(self) -> None:
        base = {
            "issue_id": "C3.1",
            "risk_level": "low",
            "issue_comment": "字段未填写。",
            "anchors": [],
            "reasoning": "缺少内容。",
            "criterion": "信息应填写完整。",
            "check_point": "信息应填写完整。",
        }

        with self.assertRaisesRegex(ValueError, "unique issue_id"):
            SubAgentOutput.model_validate(
                {"status": "issues_found", "issues": [base, dict(base)]}
            )

    def test_wrong_model_anchor_id_is_repaired_from_unique_retrieval_evidence(self) -> None:
        context = (
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p8\n内容：\n合同编号：A\n\n---\n\n"
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p15\n内容：\n甲方：浙江农村商业联合银行股份有限公司"
        )
        output = {
            "issues": [{
                "anchors": [{
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p8",
                    "quoted_text": "甲方：浙江农村商业联合银行股份有限公司",
                    "comment_text": "测试",
                }],
            }],
        }

        stats = OrchestratorAgent()._validate_and_repair_anchors(output, context)

        self.assertEqual(stats, {"valid": 0, "repaired": 1, "invalid": 0})
        self.assertEqual(output["issues"][0]["anchors"][0]["xml_anchor_id"], "path:body/p15")

    def test_exact_quote_match_wins_over_short_containment_candidate(self) -> None:
        context = (
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p4\n"
            "内容：\n深信服科技股份有限公司\n\n---\n\n"
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p19\n"
            "内容：\n甲方：浙江农村商业联合银行股份有限公司\n\n---\n\n"
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p31\n"
            "内容：\n乙方：深信服科技股份有限公司 地址：【未填写】"
        )
        output = {
            "issues": [{
                "anchors": [{
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p19",
                    "quoted_text": "乙方：深信服科技股份有限公司 地址：【未填写】",
                    "comment_text": "请填写乙方地址。",
                }],
            }],
        }

        stats = OrchestratorAgent()._validate_and_repair_anchors(output, context)

        self.assertEqual(stats, {"valid": 0, "repaired": 1, "invalid": 0})
        self.assertEqual(output["issues"][0]["anchors"][0]["xml_anchor_id"], "path:body/p31")

    def test_unprovable_model_anchor_is_preserved_for_regeneration_feedback(self) -> None:
        context = "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p1\n内容：\n合同正文"
        output = {
            "issues": [{
                "anchors": [{
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "模型编造的原文",
                    "comment_text": "测试",
                }],
            }],
        }

        stats = OrchestratorAgent()._validate_and_repair_anchors(output, context)

        self.assertEqual(stats, {"valid": 0, "repaired": 0, "invalid": 1})
        self.assertEqual(len(output["issues"][0]["anchors"]), 1)
        errors = OrchestratorAgent()._anchor_validation_errors(output, context)
        self.assertTrue(any("quoted_text 不属于锚点" in error for error in errors))

    def test_overlapping_nonidentical_edit_ranges_require_regeneration(self) -> None:
        context = (
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p1\n内容：\n"
            "甲方：某银行 乙方：某科技公司"
        )
        output = {
            "issues": [
                {
                    "issue_id": "H10.1",
                    "anchors": [{
                        "role": "edit",
                        "xml_anchor_type": "paragraph",
                        "xml_anchor_id": "path:body/p1",
                        "quoted_text": "甲方：某银行 乙方：某科技公司",
                    }],
                },
                {
                    "issue_id": "H10.2",
                    "anchors": [{
                        "role": "edit",
                        "xml_anchor_type": "paragraph",
                        "xml_anchor_id": "path:body/p1",
                        "quoted_text": "乙方：某科技公司",
                    }],
                },
            ],
        }

        errors = OrchestratorAgent()._anchor_validation_errors(output, context)

        self.assertTrue(any("edit 范围" in error and "重叠" in error for error in errors))

    def test_identical_edit_ranges_are_left_for_semantic_review(self) -> None:
        context = "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p1\n内容：\n乙方：某科技公司"
        anchor = {
            "role": "edit",
            "xml_anchor_type": "paragraph",
            "xml_anchor_id": "path:body/p1",
            "quoted_text": "乙方：某科技公司",
        }
        output = {
            "issues": [
                {"issue_id": "H10.1", "anchors": [dict(anchor)]},
                {"issue_id": "H10.2", "anchors": [dict(anchor)]},
            ],
        }

        self.assertEqual(OrchestratorAgent()._anchor_validation_errors(output, context), [])

    def test_invalid_anchor_requests_complete_model_revision(self) -> None:
        context = "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p1\n内容：\n乙方：某科技公司"
        invalid = {
            "status": "issues_found",
            "issues": [{
                "issue_id": "H10.1",
                "anchors": [{
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "模型编造的原文",
                }],
            }],
        }
        corrected = {
            "status": "issues_found",
            "issues": [{
                "issue_id": "H10.1",
                "anchors": [{
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "乙方：某科技公司",
                }],
            }],
        }

        class RevisingOrchestrator(OrchestratorAgent):
            def __init__(self):
                super().__init__()
                self.revision_prompts: list[str] = []

            async def _revise_subagent_output(self, _sub_agent, prompt, _previous):
                self.revision_prompts.append(prompt)
                return corrected, "corrected"

        orchestrator = RevisingOrchestrator()
        sub_agent = type("SubAgentStub", (), {"messages": []})()
        result, raw, _, _ = asyncio.run(
            orchestrator._validate_and_revise_output(
                sub_agent,
                invalid,
                "invalid",
                {},
                context,
                coverage_audit=False,
            )
        )

        self.assertEqual(result, corrected)
        self.assertEqual(raw, "corrected")
        self.assertEqual(len(orchestrator.revision_prompts), 1)
        self.assertIn("quoted_text 不属于锚点", orchestrator.revision_prompts[0])

    def test_retrieval_context_omits_display_ordinals(self) -> None:
        index = {
            "anchor_map": {
                "anchor-31": {
                    "body_child_index": 31,
                    "type": "p",
                    "raw_text": "乙方：深信服科技股份有限公司 地址：【未填写】",
                }
            }
        }
        nodes = [{
            "node_id": "frontmatter/node-1",
            "title": "合同首部",
            "start_index": 31,
            "end_index": 31,
        }]

        rendered = _format_anchor_context(index, nodes)

        self.assertIn("## 合同检索证据", rendered)
        self.assertNotIn("合同检索结果 1", rendered)
        self.assertIn("xml_anchor_id: path:body/p31", rendered)


class MissingTextReviewTests(unittest.IsolatedAsyncioTestCase):
    async def test_supplementary_search_uses_fixed_hybrid_route(self) -> None:
        class FakeMcpClient:
            def __init__(self) -> None:
                self.payload = None

            async def call_tool(self, name, payload):
                self.payload = payload
                return "未找到相关内容。"

        client = FakeMcpClient()
        orchestrator = OrchestratorAgent(
            mcp_client=client,
            api_events_path="events.jsonl",
        )

        await orchestrator._build_missing_text_review_notes(
            {
                "issues": [
                    {
                        "issue_id": "C5.1",
                        "check_point": "应约定保密条款",
                        "criterion": "合同条款应完整",
                        "anchors": [],
                    }
                ]
            },
            "C5",
        )

        self.assertEqual(client.payload["criterion_id"], "C5")
        self.assertEqual(
            client.payload["route_override"],
            {
                "steps": [
                    {
                        "method": "hybrid",
                        "terms": [],
                        "regions": [],
                        "all_titles": False,
                        "anchor_terms": [],
                        "coverage": "relevant",
                    }
                ],
                "fallback": False,
            },
        )


class DocxAnchorRegressionTests(unittest.TestCase):
    def setUp(self):
        patcher = patch("tools.document.reporting.docx_report.LOGGING_DETAILED", True)
        patcher.start()
        self.addCleanup(patcher.stop)

    def test_new_summary_replaces_old_ai_summary(self) -> None:
        with TemporaryDirectory() as temp_dir:
            source = Path(temp_dir) / "source.docx"
            first = Path(temp_dir) / "first.docx"
            second = Path(temp_dir) / "second.docx"
            document = Document()
            document.add_paragraph("合同正文")
            document.save(source)
            results = [{
                "criterion_id": "TEST",
                "criterion": "测试",
                "issues": [{
                    "issue_id": "TEST.1",
                    "risk_level": "low",
                    "issue_comment": "测试问题",
                    "check_point": "测试",
                    "anchors": [{
                        "role": "edit",
                        "xml_anchor_type": "paragraph",
                        "xml_anchor_id": "path:body/p1",
                        "quoted_text": "合同正文",
                        "comment_text": "测试批注",
                    }],
                }],
            }]
            DocxReportGenerator._generate_docx_with_comments(
                str(source), results, str(first), {"overall_comment": "旧总结", "priority_comments": []}
            )
            DocxReportGenerator._generate_docx_with_comments(
                str(first), results, str(second), {"overall_comment": "新总结", "priority_comments": []}
            )

            namespace = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
            with zipfile.ZipFile(second) as archive:
                comments = ET.fromstring(archive.read("word/comments.xml"))
            summaries = [
                "".join(comment.itertext())
                for comment in comments
                if comment.attrib.get(f"{{{namespace}}}author") == "AI 审查总结"
            ]
            self.assertEqual(summaries, ["总体审查结论：新总结"])

    def test_anchor_paths_keep_real_xml_child_indexes(self) -> None:
        with TemporaryDirectory() as temp_dir:
            contract_path = Path(temp_dir) / "contract.docx"
            document = Document()
            first = document.add_paragraph("第一段")
            document.add_paragraph("目标条款")
            first._p.addnext(OxmlElement("w:sdt"))
            document.save(contract_path)

            nodes = build_docx_anchor_nodes(str(contract_path))
            target = next(node for node in nodes if node.text == "目标条款")
            self.assertEqual(target.anchor_id, "path:body/p3")

            anchor = DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                Document(contract_path),
                "paragraph",
                target.anchor_id,
                "目标条款",
            )
            self.assertIsNotNone(anchor)
            self.assertEqual(anchor.path, "body/p3")

    def test_anchor_matching_tolerates_spacing_variants_within_same_anchor(self) -> None:
        with TemporaryDirectory() as temp_dir:
            contract_path = Path(temp_dir) / "contract.docx"
            document = Document()
            document.add_paragraph("甲方应在十\u00a0个工作日内付款")
            document.save(contract_path)

            anchor = DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                Document(contract_path),
                "paragraph",
                "path:body/p1",
                "甲方应在十 个工作日内付款",
            )
            self.assertIsNotNone(anchor)
            self.assertEqual(anchor.match.matched_text, "甲方应在十\u00a0个工作日内付款")

    def test_table_row_quote_anchors_complete_contained_cell_paragraph(self) -> None:
        with TemporaryDirectory() as temp_dir:
            contract_path = Path(temp_dir) / "contract.docx"
            document = Document()
            table = document.add_table(rows=1, cols=2)
            table.cell(0, 0).text = "备 注"
            table.cell(0, 1).text = "设备运到招标单位指定地点。"
            document.save(contract_path)

            anchor = DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                Document(contract_path),
                "table",
                "path:body/tbl1",
                "备 注\t设备运到招标单位指定地点。",
            )

            self.assertIsNotNone(anchor)
            self.assertEqual(anchor.path, "body/tbl1/r1c2/p2")
            self.assertEqual(anchor.match.matched_text, "设备运到招标单位指定地点。")

    def test_anchorless_issue_is_written_and_diagnosed(self) -> None:
        with TemporaryDirectory() as temp_dir:
            contract_path = Path(temp_dir) / "contract.docx"
            output_path = Path(temp_dir) / "reviewed.docx"
            annotation_path = Path(temp_dir) / "logs" / "annotation_events.json"
            document = Document()
            document.add_paragraph("合同正文")
            document.add_paragraph("（以下无正文）")
            document.save(contract_path)

            result = DocxReportGenerator.generate_annotated_docx(
                str(contract_path),
                [{
                    "criterion_id": "TEST",
                    "criterion": "正文最后应有空白声明",
                    "issues": [{
                        "issue_id": "TEST.1",
                        "risk_level": "low",
                        "issue_comment": "建议在正文最后补充空白声明。",
                        "check_point": "正文最后",
                        "anchors": [],
                    }],
                }],
                str(output_path),
                annotation_events_path=str(annotation_path),
                region_boundaries={
                    "body_content_end": {
                        "xml_anchor_type": "paragraph",
                        "xml_anchor_id": "path:body/p1",
                        "quoted_text": "合同正文",
                    }
                },
            )

            self.assertEqual(result, str(output_path))
            report = json.loads(annotation_path.read_text(encoding="utf-8"))
            self.assertEqual(report["stats"]["missing_text_fallback"], 1)
            self.assertEqual(report["skipped_count"], 0)
            self.assertEqual(report["events"][0]["status"], "anchored")
            self.assertEqual(report["events"][0]["match_strategy"], "body_content_end_missing_issue")

    def test_missing_body_boundary_does_not_fall_back_to_attachment_end(self) -> None:
        with TemporaryDirectory() as temp_dir:
            contract_path = Path(temp_dir) / "contract.docx"
            output_path = Path(temp_dir) / "reviewed.docx"
            annotation_path = Path(temp_dir) / "logs" / "annotation_events.json"
            document = Document()
            document.add_paragraph("合同正文")
            document.add_paragraph("附件中的最后一段")
            document.save(contract_path)

            DocxReportGenerator.generate_annotated_docx(
                str(contract_path),
                [{
                    "criterion_id": "TEST",
                    "criterion": "正文结尾应有空白声明",
                    "issues": [{
                        "issue_id": "TEST.1",
                        "risk_level": "low",
                        "issue_comment": "建议在正文结尾补充空白声明。",
                        "check_point": "正文结尾",
                        "anchors": [],
                    }],
                }],
                str(output_path),
                annotation_events_path=str(annotation_path),
            )

            report = json.loads(annotation_path.read_text(encoding="utf-8"))
            self.assertEqual(report["skipped_count"], 1)
            self.assertEqual(report["events"][0]["status"], "unmatched")
            self.assertEqual(report["events"][0]["match_strategy"], "body_content_end_unavailable")

    def test_summary_comment_includes_complete_summary_only_issues(self) -> None:
        text = DocxReportGenerator._build_summary_comment_text({
            "overall_comment": "总体概述。",
            "priority_comments": [],
            "opening_notices": ["注意：补充固定提示语。"],
        })

        self.assertIn("合同首部需补充的提示：", text)
        self.assertIn("注意：补充固定提示语。", text)
        self.assertNotIn("全局审查问题：", text)

    def test_summary_comment_discloses_incomplete_reviews_without_internal_id(self) -> None:
        results = [{
            "criterion_id": "H11",
            "section": "合同框架",
            "criterion": "校验正文与附件约定是否一致。",
            "status": "ERROR",
        }]

        notices = _incomplete_review_notices(results)
        text = DocxReportGenerator._build_summary_comment_text({
            "overall_comment": "其余审查项已完成。",
            "priority_comments": [],
            "incomplete_reviews": notices,
        })

        self.assertIn("未完成审查项：", text)
        self.assertIn("校验正文与附件约定是否一致", text)
        self.assertNotIn("H11", text)


class PromptAuthorityTests(unittest.TestCase):
    def test_explicit_criteria_are_authoritative_for_review_and_reflection(self) -> None:
        prompt_path = Path(__file__).resolve().parents[2] / "agents" / "prompts" / "cn_prompts.yaml"

        prompts = load_prompts(prompt_path)

        self.assertIn("先判断当前标准的具体检查对象与触发条件", prompts["SUB_AGENT_BASE_PROMPT"])
        self.assertIn("独立判断事实是否真正触发当前标准", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertIn("指定提示语", prompts["SUB_AGENT_BASE_PROMPT"])
        self.assertIn("固定文字完整性", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertNotIn("output_scope 枚举", prompts["SUMMARIZER_SYSTEM_PROMPT"])
        self.assertIn("不得输出 issue_decisions", prompts["SUMMARIZER_SYSTEM_PROMPT"])
        self.assertIn("不得使用“断号”", prompts["SUMMARIZER_SYSTEM_PROMPT"])
        self.assertIn("付款计划数量必须按", prompts["SUB_AGENT_BASE_PROMPT"])
        self.assertIn("付款节点计数", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertIn("检查合同中所有未填写内容", prompts["PLANNER_STRATEGY_PROMPT"])
        self.assertIn("不得按固定字符形态机械判断", prompts["SUB_AGENT_BASE_PROMPT"])
        self.assertIn("列举项不是封闭清单", prompts["SUB_AGENT_BASE_PROMPT"])
        self.assertIn("不得按固定正则形态判断", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertIn("每个 issue 最多只能对应一个 edit", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertIn("不得让不同 issue 重复引用整段", prompts["SUB_AGENT_BASE_PROMPT"])
        self.assertIn("没有角色标签", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertIn("必须列入 coverage_gaps", prompts["REFLECTOR_SYSTEM_PROMPT"])

    def test_runtime_prompt_makes_deterministic_absence_authoritative(self) -> None:
        source = inspect.getsource(OrchestratorAgent.execute_single_criterion)

        self.assertIn("确定性计数和范围说明", source)
        self.assertIn("必须输出 issues_found，不得输出 compliant", source)


class SubAgentSchemaNormalizationTests(unittest.TestCase):
    def test_word_comments_preserve_complete_text_within_new_limits(self) -> None:
        output = SubAgentOutput.model_validate({
            "status": "issues_found",
            "issues": [{
                "issue_id": "C11.1",
                "risk_level": "medium",
                "issue_comment": "问" * 180,
                "anchors": [{
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "合同原文",
                    "comment_text": "改" * 120,
                }],
                "reasoning": "正文与附件日期不一致",
                "criterion": "检查正文与附件一致性",
                "check_point": "日期必须一致",
            }],
        })

        self.assertEqual(len(output.issues[0].issue_comment), 180)
        self.assertEqual(len(output.issues[0].anchors[0].comment_text), 120)


class WorkflowPresentationTests(unittest.TestCase):
    def setUp(self) -> None:
        self.results = [
            {
                "criterion_id": "C12",
                "section": "文字",
                "criterion": "检查重复词",
                "status": "issues_found",
                "issues": [{
                    "issue_id": "C12.1",
                    "risk_level": "low",
                    "issue_comment": "存在重复用词。",
                    "anchors": [{
                        "xml_anchor_type": "paragraph",
                        "xml_anchor_id": "path:body/p1",
                        "quoted_text": "因乙方原因原因",
                        "comment_text": "删除重复文字。",
                    }],
                    "reasoning": "原因重复",
                    "criterion": "检查重复词",
                    "check_point": "重复词",
                }],
            },
            {
                "criterion_id": "C15",
                "section": "语法",
                "criterion": "检查语病",
                "status": "issues_found",
                "issues": [{
                    "issue_id": "C15.1",
                    "risk_level": "low",
                    "issue_comment": "存在语病。",
                    "anchors": [{
                        "xml_anchor_type": "paragraph",
                        "xml_anchor_id": "path:body/p1",
                        "quoted_text": "因乙方原因原因",
                        "comment_text": "修改语病。",
                    }],
                    "reasoning": "原因重复",
                    "criterion": "检查语病",
                    "check_point": "语病",
                }],
            },
        ]

    def test_summary_schema_rejects_unknown_issue_decisions_field(self) -> None:
        with self.assertRaises(ValueError):
            SummaryOutput.model_validate({
                "overall_comment": "总结",
                "priority_comments": [],
                "issue_decisions": [],
            })

    def test_workflow_does_not_preempt_summarizer_semantic_deduplication(self) -> None:
        duplicate = json.loads(json.dumps(self.results[0], ensure_ascii=False))
        duplicate["criterion_id"] = "C18"
        duplicate["issues"][0]["anchors"][0]["comment_text"] = "请填写该字段。"

        inline, _ = _prepare_inline_results([self.results[0], duplicate])

        self.assertEqual(len(inline), 2)

    def test_pending_criterion_keeps_confirmed_findings_only(self) -> None:
        opinion = {"status": "issues_found", "issues": self.results[0]["issues"]}
        ledger = _FindingLedger(opinion)
        finding_id = ledger.current(opinion)[0]["finding_id"]
        ledger.review(
            opinion,
            {
                "status": "REJECT",
                "feedback": "其他 finding 仍需修改",
                "finding_decisions": [
                    {"finding_id": finding_id, "decision": "confirm", "reason": "事实成立"}
                ],
            },
            1,
        )

        confirmed = ledger.confirmed_issues(opinion)

        self.assertEqual(len(confirmed), 1)
        self.assertEqual(confirmed[0]["finding_id"], finding_id)

    def test_partial_reflector_decisions_are_not_discarded(self) -> None:
        opinion = {
            "status": "issues_found",
            "issues": [
                self.results[0]["issues"][0],
                self.results[1]["issues"][0],
            ],
        }
        ledger = _FindingLedger(opinion)
        first_id = ledger.current(opinion)[0]["finding_id"]

        passed = ledger.review(
            opinion,
            {
                "status": "REJECT",
                "feedback": "另一项仍需判断",
                "finding_decisions": [
                    {"finding_id": first_id, "decision": "confirm", "reason": "事实成立"}
                ],
            },
            1,
        )

        self.assertFalse(passed)
        self.assertEqual(
            [issue["finding_id"] for issue in ledger.confirmed_issues(opinion)],
            [first_id],
        )
        self.assertEqual(len(ledger.unresolved_issues(opinion)), 1)

    def test_confirmed_finding_survives_later_model_omission(self) -> None:
        opinion = {"status": "issues_found", "issues": self.results[0]["issues"]}
        ledger = _FindingLedger(opinion)
        finding_id = ledger.current(opinion)[0]["finding_id"]
        ledger.review(
            opinion,
            {
                "status": "PASS",
                "feedback": "",
                "finding_decisions": [
                    {"finding_id": finding_id, "decision": "confirm", "reason": "事实成立"}
                ],
            },
            1,
        )

        confirmed = ledger.confirmed_issues({"status": "compliant", "issues": []})

        self.assertEqual(len(confirmed), 1)
        self.assertEqual(confirmed[0]["finding_id"], finding_id)

    def test_finding_id_stays_stable_when_comment_and_anchor_are_revised(self) -> None:
        opinion = {"status": "issues_found", "issues": self.results[0]["issues"]}
        ledger = _FindingLedger(opinion)
        original_id = ledger.current(opinion)[0]["finding_id"]
        revised = json.loads(json.dumps(opinion, ensure_ascii=False))
        revised["issues"][0]["issue_comment"] = "修正后的独立问题说明。"
        revised["issues"][0]["anchors"] = []
        ledger.observe(revised)

        self.assertEqual(ledger.current(revised)[0]["finding_id"], original_id)
        self.assertEqual(ledger.missing(revised), [])

    def test_comparison_policy_requires_edit_and_evidence(self) -> None:
        opinion = {"status": "issues_found", "issues": self.results[0]["issues"]}
        context = (
            "source_region: attachments\n"
            "xml_anchor_type: paragraph\n"
            "xml_anchor_id: path:body/p200\n"
            "内容：\n附件对照原文。"
        )
        self.assertEqual(
            len(OrchestratorAgent._policy_output_errors(
                opinion,
                {"evidence_region": "attachments"},
                context,
            )),
            1,
        )
        opinion["issues"][0]["anchors"].append({
            "role": "evidence",
            "xml_anchor_type": "paragraph",
            "xml_anchor_id": "path:body/p200",
            "quoted_text": "附件对照原文。",
            "comment_text": "",
        })

        self.assertEqual(
            OrchestratorAgent._policy_output_errors(
                opinion,
                {"evidence_region": "attachments"},
                context,
            ),
            [],
        )

    def test_policy_rejects_reversed_edit_and_evidence_regions(self) -> None:
        context = (
            "source_region: frontmatter\n"
            "xml_anchor_type: paragraph\n"
            "xml_anchor_id: path:body/p25\n"
            "内容：\n乙方：首部公司\n\n---\n\n"
            "source_region: tail\n"
            "xml_anchor_type: paragraph\n"
            "xml_anchor_id: path:body/p111\n"
            "内容：\n乙方：签署公司"
        )
        issue = {
            "issue_id": "H10.1",
            "anchors": [
                {
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p25",
                },
                {
                    "role": "evidence",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p111",
                },
            ],
        }
        policy = {"edit_region": "tail", "evidence_region": "frontmatter"}

        errors = OrchestratorAgent._policy_output_errors(
            {"status": "issues_found", "issues": [issue]},
            policy,
            context,
        )

        self.assertEqual(len(errors), 2)
        self.assertTrue(any("edit 不来自 tail region" in error for error in errors))
        self.assertTrue(any("evidence 不来自 frontmatter region" in error for error in errors))

    def test_policy_accepts_tail_edit_and_frontmatter_evidence(self) -> None:
        context = (
            "source_region: frontmatter\n"
            "xml_anchor_type: paragraph\n"
            "xml_anchor_id: path:body/p25\n"
            "内容：\n乙方：首部公司\n\n---\n\n"
            "source_region: tail\n"
            "xml_anchor_type: paragraph\n"
            "xml_anchor_id: path:body/p111\n"
            "内容：\n乙方：签署公司"
        )
        issue = {
            "issue_id": "H10.1",
            "anchors": [
                {
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p111",
                },
                {
                    "role": "evidence",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p25",
                },
            ],
        }

        self.assertEqual(
            OrchestratorAgent._policy_output_errors(
                {"status": "issues_found", "issues": [issue]},
                {"edit_region": "tail", "evidence_region": "frontmatter"},
                context,
            ),
            [],
        )

    def test_ambiguous_different_advice_is_not_merged_only_by_criterion(self) -> None:
        for result in self.results:
            result["issues"][0]["anchors"][0]["quoted_text"] = "乙方向甲方发贷。"
        self.results[0]["issues"][0]["anchors"][0]["comment_text"] = (
            "若意为交付货物请改为“发货”，若为无关内容请删除。"
        )
        self.results[1]["issues"][0]["anchors"][0]["comment_text"] = (
            "结合语境应改为“乙方向甲方发货”。"
        )

        inline, _ = _prepare_inline_results(self.results)
        self.assertEqual(len(inline), 2)

    def test_summary_duplicate_group_removes_semantic_duplicate_issue(self) -> None:
        first = json.loads(json.dumps(self.results[0], ensure_ascii=False))
        second = json.loads(json.dumps(self.results[0], ensure_ascii=False))
        first["criterion_id"] = "H12"
        second["criterion_id"] = "H15"
        second["issues"][0]["issue_id"] = "H15.1"

        merged, stats = apply_summary_duplicate_groups(
            [first, second],
            [{
                "keep_issue_key": "0:H12:C12.1",
                "merge_issue_keys": ["1:H15:H15.1"],
                "canonical_comment": "删除重复的“原因”，改为“因乙方原因”。",
                "reason": "同一处重复词问题",
            }],
        )

        self.assertEqual([len(result["issues"]) for result in merged], [1, 0])
        self.assertEqual(stats["removed_duplicate_issues"], 1)
        self.assertEqual(
            merged[0]["issues"][0]["anchors"][0]["comment_text"],
            "删除重复的“原因”，改为“因乙方原因”。",
        )

    def test_summary_duplicate_group_keeps_issues_at_different_locations(self) -> None:
        first = json.loads(json.dumps(self.results[0], ensure_ascii=False))
        second = json.loads(json.dumps(self.results[0], ensure_ascii=False))
        first["criterion_id"] = "H01"
        second["criterion_id"] = "H01"
        second["issues"][0]["issue_id"] = "C12.2"
        second["issues"][0]["anchors"][0]["xml_anchor_id"] = "path:body/p31"
        second["issues"][0]["anchors"][0]["quoted_text"] = "乙方地址：【未填写】"

        merged, stats = apply_summary_duplicate_groups(
            [first, second],
            [{
                "keep_issue_key": "0:H01:C12.1",
                "merge_issue_keys": ["1:H01:C12.2"],
                "canonical_comment": "合同首部信息未填写。",
                "reason": "均为未填写",
            }],
        )

        self.assertEqual([len(result["issues"]) for result in merged], [1, 1])
        self.assertEqual(stats["applied_groups"], 0)
        self.assertEqual(stats["skipped_groups"], 1)
        self.assertEqual(stats["removed_duplicate_issues"], 0)

    def test_summary_issue_cannot_absorb_anchored_inline_issue(self) -> None:
        inline = json.loads(json.dumps(self.results[0], ensure_ascii=False))
        inline["criterion_id"] = "H01"
        inline["issues"][0]["issue_id"] = "1.4"
        inline["issues"][0]["anchors"][0]["xml_anchor_id"] = "path:body/p31"
        inline["issues"][0]["anchors"][0]["quoted_text"] = (
            "乙方：深信服科技股份有限公司 地址：【未填写】"
        )
        summary = json.loads(json.dumps(inline, ensure_ascii=False))
        summary["issues"][0]["issue_id"] = "1.5"
        summary["issues"][0]["anchors"] = []

        merged, stats = apply_summary_duplicate_groups(
            [inline, summary],
            [{
                "keep_issue_key": "1:H01:1.5",
                "merge_issue_keys": ["0:H01:1.4"],
                "canonical_comment": "合同首部存在多处信息未填写。",
                "reason": "均为未填写",
            }],
        )

        self.assertEqual([len(result["issues"]) for result in merged], [1, 1])
        self.assertEqual(stats["applied_groups"], 0)
        self.assertEqual(stats["skipped_groups"], 1)
        self.assertEqual(stats["removed_duplicate_issues"], 0)

    def test_keyword_triggered_fixed_instruction_is_routed_to_opening_summary(self) -> None:
        results = [{
            "criterion_id": "FIXED",
            "criterion": "如果标题出现采购，就直接在合同首部写明“注意：完整固定文字”。",
            "status": "issues_found",
            "issues": [{
                "issue_id": "FIXED.1",
                "risk_level": "low",
                "issue_comment": "缺少提示语。",
                "anchors": [{
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p1",
                    "quoted_text": "采购合同",
                    "comment_text": "建议补充提示语。",
                }],
                "reasoning": "标准触发",
                "criterion": "如果标题出现采购，就直接在合同首部写明“注意：完整固定文字”。",
                "check_point": "直接写明“注意：当前检查点固定文字”",
            }],
        }]

        results[0]["placement"] = {"target": "summary"}
        results[0]["review_instruction"] = "注意：当前检查点固定文字"
        inline, notices = _prepare_inline_results(results)
        self.assertEqual(inline, [])
        self.assertEqual(notices, ["注意：当前检查点固定文字"])

    def test_fixed_instruction_requires_confirmed_issue_status(self) -> None:
        base = {
            "criterion_id": "FIXED",
            "criterion": "触发后输出固定提示。",
            "placement": {"target": "summary"},
            "review_instruction": "注意：固定提示。",
        }
        confirmed_issue = json.loads(json.dumps(self.results[0]["issues"][0], ensure_ascii=False))

        cases = [
            ({**base, "status": "NOT_APPLICABLE", "issues": []}, []),
            ({**base, "status": "COMPLIANT", "issues": []}, []),
            ({**base, "status": "ISSUES_FOUND", "issues": []}, []),
            ({**base, "status": "ISSUES_FOUND", "issues": [confirmed_issue]}, ["注意：固定提示。"]),
        ]

        for result, expected_notices in cases:
            with self.subTest(status=result["status"], issue_count=len(result["issues"])):
                inline, notices = _prepare_inline_results([result])
                self.assertEqual(inline, [])
                self.assertEqual(notices, expected_notices)

    def test_summary_placement_remains_authoritative_for_missing_fields(self) -> None:
        result = self.results[0]
        result["placement"] = {"target": "summary"}
        result["review_instruction"] = "注意：请补充合同基本信息"

        inline, notices = _prepare_inline_results([result])

        self.assertEqual(inline, [])
        self.assertEqual(notices, ["注意：请补充合同基本信息"])

    def test_user_facing_comments_reject_internal_locations_and_same_as_above(self) -> None:
        value = json.loads(json.dumps(self.results[0]["issues"][0], ensure_ascii=False))
        value["issue_comment"] = "请与p311保持一致。"
        with self.assertRaises(ValueError):
            SubAgentOutput.model_validate({"status": "issues_found", "issues": [value]})

        value["issue_comment"] = "存在文字错误。"
        value["anchors"][0]["comment_text"] = "同上，请修正。"
        with self.assertRaises(ValueError):
            SubAgentOutput.model_validate({"status": "issues_found", "issues": [value]})

    def test_complete_region_coverage_requires_each_evidence_anchor(self) -> None:
        orchestrator = OrchestratorAgent()
        context = (
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p8\n内容：\n合同号：[待填写]\n\n---\n\n"
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p9\n内容：\n签订日期：2025年 月 日"
        )
        output = {
            "issues": [],
            "coverage_decisions": [{
                "xml_anchor_type": "paragraph",
                "xml_anchor_id": "path:body/p8",
                "decision": "compliant",
                "reason": "已审查",
            }],
        }

        errors = orchestrator._coverage_audit_errors(output, context)

        self.assertEqual(len(errors), 1)
        self.assertIn("path:body/p9", errors[0])

    def test_complete_region_coverage_rejects_issue_anchor_marked_compliant(self) -> None:
        orchestrator = OrchestratorAgent()
        context = (
            "xml_anchor_type: paragraph\nxml_anchor_id: path:body/p31\n"
            "内容：\n乙方：深信服科技股份有限公司 地址：【未填写】"
        )
        output = {
            "issues": [{
                "anchors": [{
                    "role": "edit",
                    "xml_anchor_type": "paragraph",
                    "xml_anchor_id": "path:body/p31",
                    "quoted_text": "乙方：深信服科技股份有限公司 地址：【未填写】",
                    "comment_text": "请填写乙方地址。",
                }],
            }],
            "coverage_decisions": [{
                "xml_anchor_type": "paragraph",
                "xml_anchor_id": "path:body/p31",
                "decision": "compliant",
                "reason": "已审查",
            }],
        }

        errors = orchestrator._coverage_audit_errors(output, context)

        self.assertEqual(len(errors), 1)
        self.assertIn("未标记为 issue", errors[0])
        self.assertIn("path:body/p31", errors[0])


class CriterionFailureIsolationTests(unittest.IsolatedAsyncioTestCase):
    @staticmethod
    def _criteria() -> list[dict]:
        return [
            {"id": "H01", "section": "合同主体", "criterion": "检查合同主体信息。"},
            {"id": "H11", "section": "合同框架", "criterion": "校验正文与附件约定是否一致。"},
        ]

    async def test_one_schema_failure_keeps_other_criterion_result(self) -> None:
        orchestrator = OrchestratorAgent()

        async def execute(criterion: dict, *, max_attempts: int) -> dict:
            self.assertEqual(max_attempts, 2)
            if criterion["id"] == "H11":
                raise AgentOutputValidationError("invalid user-facing output")
            return {
                "criterion_id": criterion["id"],
                "criterion": criterion["criterion"],
                "section": criterion["section"],
                "issues": [],
                "status": "COMPLIANT",
                "quality_status": "not_required",
                "tokens": 0,
            }

        orchestrator._execute_criterion_with_recovery = execute  # type: ignore[method-assign]

        results = await orchestrator.execute_criteria(self._criteria())

        self.assertEqual([result["status"] for result in results], ["COMPLIANT", "ERROR"])
        self.assertEqual(results[1]["quality_status"], "error")
        self.assertEqual(results[1]["error_type"], "AgentOutputValidationError")

    async def test_all_schema_failures_remain_task_level_failure(self) -> None:
        orchestrator = OrchestratorAgent()

        async def execute(_criterion: dict, *, max_attempts: int) -> dict:
            self.assertEqual(max_attempts, 2)
            raise AgentOutputValidationError("invalid user-facing output")

        orchestrator._execute_criterion_with_recovery = execute  # type: ignore[method-assign]

        with self.assertRaisesRegex(RuntimeError, "All criteria failed"):
            await orchestrator.execute_criteria(self._criteria())

    async def test_unexpected_single_criterion_error_keeps_other_result(self) -> None:
        orchestrator = OrchestratorAgent()

        async def execute(criterion: dict, *, max_attempts: int) -> dict:
            self.assertEqual(max_attempts, 2)
            if criterion["id"] == "H01":
                raise ValueError("post-revision validation failed")
            return {
                "criterion_id": criterion["id"],
                "criterion": criterion["criterion"],
                "section": criterion["section"],
                "issues": [],
                "status": "COMPLIANT",
                "quality_status": "not_required",
                "tokens": 0,
            }

        orchestrator._execute_criterion_with_recovery = execute  # type: ignore[method-assign]

        results = await orchestrator.execute_criteria(self._criteria())

        self.assertEqual([result["status"] for result in results], ["ERROR", "COMPLIANT"])
        self.assertEqual(results[0]["error_type"], "ValueError")


class SummarizerPromptBudgetTests(unittest.TestCase):
    def test_summary_output_replaces_internal_numbering_jargon(self) -> None:
        summary = normalize_summary_output({
            "overall_comment": "合同条款断号严重。",
            "priority_comments": ["修正序号断号。"],
        })

        self.assertEqual(summary["overall_comment"], "合同条款编号多处不连续。")
        self.assertEqual(summary["priority_comments"], ["修正编号不连续。"])

    def test_conservative_bound_accepts_small_prompt(self) -> None:
        bound = conservative_message_token_upper_bound([
            {"role": "system", "content": "总结合同"},
            {"role": "user", "content": "问题列表"},
        ])

        self.assertLess(bound, SUMMARIZER_MAX_PROMPT_TOKENS)

    def test_conservative_bound_rejects_prompt_above_hard_limit(self) -> None:
        messages = [
            {"role": "user", "content": " x" * (SUMMARIZER_MAX_PROMPT_TOKENS + 1)},
        ]
        bound = conservative_message_token_upper_bound(messages)

        self.assertGreater(bound, SUMMARIZER_MAX_PROMPT_TOKENS)
        summarizer = SummarizerAgent.__new__(SummarizerAgent)
        summarizer.name = "Summarizer"
        summarizer.messages = messages
        summarizer.tools = []
        summarizer.MAX_CONTEXT_TOKENS = 254_000
        summarizer.MAX_PROMPT_TOKENS = SUMMARIZER_MAX_PROMPT_TOKENS
        summarizer.debug = False
        with self.assertRaisesRegex(ValueError, "per-call safety limit"):
            summarizer._check_context_limits()

    def test_global_prompt_budget_includes_tool_schema(self) -> None:
        messages = [{"role": "user", "content": "检查合同"}]
        tools = [{
            "type": "function",
            "function": {
                "name": "contract_search",
                "description": " x" * 1000,
                "parameters": {"type": "object", "properties": {}},
            },
        }]

        self.assertGreater(
            estimate_chat_prompt_tokens(messages, tools),
            estimate_chat_prompt_tokens(messages),
        )


class RevisionIdentityTests(unittest.TestCase):
    def test_revision_restores_issue_id_for_same_source_defect(self) -> None:
        previous = _subagent_output("1.8", "地址未填写。", "path:body/p31")
        revised = _subagent_output("1.5", "该地址字段仍未填写。", "path:body/p31")

        normalized = _preserve_revision_issue_ids(previous, revised)

        self.assertEqual(normalized["issues"][0]["issue_id"], "1.8")

    def test_revision_keeps_new_id_for_different_source_defect(self) -> None:
        previous = _subagent_output("1", "期限不明确。", "path:body/p160")
        revised = _subagent_output("2", "签订日期未填写。", "path:body/p17")

        normalized = _preserve_revision_issue_ids(previous, revised)

        self.assertEqual(normalized["issues"][0]["issue_id"], "2")

    def test_revision_assigns_collision_free_id_to_inserted_finding(self) -> None:
        previous = {
            "status": "issues_found",
            "applicability_reason": "",
            "issues": [
                _subagent_output("1.9", "账号未填写。", "path:body/p37")["issues"][0],
                {
                    **_subagent_output("1.10", "首部提示缺失。", "path:body/p40")["issues"][0],
                    "anchors": [],
                },
            ],
            "coverage_decisions": [],
        }
        revised = {
            "status": "issues_found",
            "applicability_reason": "",
            "issues": [
                _subagent_output("1.9", "开户银行未填写。", "path:body/p36")["issues"][0],
                _subagent_output("1.10", "账号未填写。", "path:body/p37")["issues"][0],
                {
                    **_subagent_output("1.11", "首部提示缺失。", "path:body/p40")["issues"][0],
                    "anchors": [],
                },
            ],
            "coverage_decisions": [],
        }

        normalized = _preserve_revision_issue_ids(previous, revised)

        ids_by_comment = {issue["issue_comment"]: issue["issue_id"] for issue in normalized["issues"]}
        self.assertEqual(ids_by_comment["账号未填写。"], "1.9")
        self.assertEqual(ids_by_comment["开户银行未填写。"], "1.9.new1")
        self.assertEqual(len(set(ids_by_comment.values())), 3)


class VisibleRevisionTextTests(unittest.TestCase):
    @staticmethod
    def _append_text(parent, text: str, *, text_tag: str = "w:t") -> None:
        run = OxmlElement("w:r")
        text_element = OxmlElement(text_tag)
        text_element.text = text
        run.append(text_element)
        parent.append(run)

    def test_index_and_report_use_same_current_revision_text(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "tracked.docx"
            document = Document()
            paragraph = document.add_paragraph()
            self._append_text(paragraph._p, "当前文字")

            deleted = OxmlElement("w:del")
            self._append_text(deleted, "已删除文字", text_tag="w:delText")
            paragraph._p.append(deleted)

            moved_from = OxmlElement("w:moveFrom")
            self._append_text(moved_from, "已移走文字")
            paragraph._p.append(moved_from)

            inserted = OxmlElement("w:ins")
            self._append_text(inserted, "新增文字")
            paragraph._p.append(inserted)
            document.save(path)

            items = read_docx_items(path)
            report_nodes = build_docx_anchor_nodes(str(path))

        paragraph_item = next(item for item in items if item.kind == "p")
        paragraph_node = next(node for node in report_nodes if node.anchor_type == "paragraph")
        self.assertEqual(paragraph_item.raw_text, "当前文字新增文字")
        self.assertEqual(paragraph_item.text, paragraph_node.text)
        self.assertNotIn("已删除文字", paragraph_item.text)
        self.assertNotIn("已移走文字", paragraph_item.text)

    def test_anchor_collectors_do_not_drop_later_paragraphs_when_object_ids_recycle(self) -> None:
        with TemporaryDirectory() as temp_dir:
            path = Path(temp_dir) / "many-paragraphs.docx"
            document = Document()
            for index in range(300):
                document.add_paragraph(f"第{index + 1}个可见段落")
            document.save(path)

            anchor_nodes = build_docx_anchor_nodes(str(path))
            report_document = Document(path)
            report_views = DocxReportGenerator._collect_paragraph_text_views(report_document)

        paragraph_nodes = [node for node in anchor_nodes if node.anchor_type == "paragraph"]
        self.assertEqual(len(paragraph_nodes), 300)
        self.assertEqual(len(report_views), 300)
        self.assertEqual(paragraph_nodes[-1].anchor_id, "path:body/p300")
        self.assertEqual(report_views[-1].path, "body/p300")


class FindingEntryIdentityTests(unittest.TestCase):
    def test_multi_edit_issue_gets_unique_local_issue_ids(self) -> None:
        issue = _subagent_output("3", "两个位置分别需要修改。", "path:body/p64")["issues"][0]
        second_anchor = {
            **issue["anchors"][0],
            "xml_anchor_id": "path:body/p65",
            "quoted_text": "第二处原文",
        }
        issue["anchors"].append(second_anchor)
        opinion = {"status": "issues_found", "issues": [issue], "coverage_decisions": []}

        entries = _FindingLedger(opinion).current(opinion)

        self.assertEqual([entry["issue_id"] for entry in entries], ["3.1", "3.2"])


class PreparedContractValidationTests(unittest.TestCase):
    def test_contract_title_selection_stays_in_early_frontmatter(self) -> None:
        values = [
            "甲方公司",
            "与",
            "乙方公司",
            "人行支付系统重构研发合同",
            "合同号：ABC",
            "正文中的合同总价",
            "合同金额及支付方式",
        ]

        selected = select_contract_title_group(
            values,
            tuple((index,) for index in range(len(values))),
            frontmatter_end=5,
        )

        self.assertEqual(selected, (3,))


class RetrievalFallbackTests(unittest.TestCase):
    def test_rule_uses_semantic_fallback_only_when_plan_requests_it(self) -> None:
        plan = RouteResponse.model_validate({
            "steps": [{"method": "rule", "terms": ["人民法院"]}],
            "fallback": True,
        })

        self.assertTrue(
            _should_use_fallback(plan, fallback_enabled=True)
        )

    def test_region_cannot_request_global_fallback(self) -> None:
        with self.assertRaises(ValueError):
            RouteResponse.model_validate({
                "steps": [{"method": "region", "regions": ["body"]}],
                "fallback": True,
            })

    def test_full_scan_does_not_duplicate_semantic_fallback(self) -> None:
        plan = RouteResponse.model_validate({
            "steps": [{"method": "scan"}],
            "fallback": False,
        })

        self.assertFalse(
            _should_use_fallback(plan, fallback_enabled=True)
        )


if __name__ == "__main__":
    unittest.main()
