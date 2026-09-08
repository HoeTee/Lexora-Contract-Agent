"""Regression tests for field coverage, evidence roles and reflection acceptance.

These verify deterministic contracts, not semantic LLM accuracy.
"""
import json
from pathlib import Path
import tempfile
import unittest
from unittest.mock import AsyncMock, patch
import xml.etree.ElementTree as ET
import zipfile

from docx import Document
from pydantic import ValidationError

from agents.orchestrator import OrchestratorAgent, SUB_AGENT_EXPECTED_JSON
from agents.reflector import REFLECTOR_EXPECTED_JSON
from agents.review_presentation import build_review_presentation
from agents.review_quality import FindingLedger
from agents.schemas import SubAgentOutput, SubAgentAnchor, ReflectorOutput, ExtractedPlacement, ExtractedRoute
from agents.prompts.cn_prompts import load_prompts, PROMPTS_PATH
from tools.document.reporting.docx_report import DocxReportGenerator


def anchor(text="地址：#地址#", path="path:body/p1", comment="地址未填写，请填写实际地址。", **extra):
    return {"role": "edit", "defect_type": "missing_value", "xml_anchor_type": "paragraph",
            "xml_anchor_id": path, "quoted_text": text, "comment_text": comment, **extra}


def opinion(anchors=None):
    return {"status": "issues_found", "issues": [{"issue_id": "1", "risk_level": "low",
            "issue_comment": "基础字段缺失", "criterion": "检查基础信息", "check_point": "检查基础信息",
            "reasoning": "原文是未替换的模板变量", "anchors": [anchor()] if anchors is None else anchors}]}


def result(value=None, **extra):
    return {"criterion_id": "C3", "criterion": "检查基础信息", "section": "主体",
            **(value or opinion()), **extra}


class QualityTests(unittest.TestCase):
    def test_expected_json_contracts(self):
        SubAgentOutput.model_validate(json.loads(SUB_AGENT_EXPECTED_JSON))
        ReflectorOutput.model_validate(json.loads(REFLECTOR_EXPECTED_JSON))

    def test_role_schema(self):
        SubAgentAnchor.model_validate(anchor(role="evidence", comment_text=""))
        for a in [anchor(role="bad"), anchor(comment_text=""), anchor(role="evidence")]:
            with self.assertRaises(ValidationError):
                SubAgentAnchor.model_validate(a)
        with self.assertRaises(ValidationError):
            ExtractedPlacement(target="merged_by_party")

    def test_long_comments_are_rejected_not_silently_truncated(self):
        with self.assertRaises(ValidationError):
            SubAgentAnchor.model_validate(anchor(comment_text="字" * 301))

    def test_evidence_only_issue_requires_edit_or_summary(self):
        with self.assertRaises(ValidationError):
            SubAgentOutput.model_validate(opinion([anchor(role="evidence", comment_text="")]))

    def test_withdrawal_does_not_permanently_veto_a_revised_finding(self):
        o = opinion()
        ledger = FindingLedger(o)
        key = ledger.current(o)[0]["finding_id"]
        removed = {"status": "compliant", "issues": []}
        self.assertTrue(ledger.review(removed, {"status": "PASS", "removals": [
            {"finding_id": key, "reason": "原判断错误"}]}, 1))
        revised = opinion([anchor(comment="请填写实际联系地址。")])
        ledger.observe(revised)
        self.assertTrue(ledger.review(revised, {"status": "PASS", "removals": []}, 2))
        self.assertNotIn(key, ledger.removals)

    def test_edit_locations_are_not_party_merged(self):
        a = [anchor(), anchor("电话：#电话#", "path:body/p2", "电话未填写。")]
        p = build_review_presentation([result(opinion(a), placement={"target": "inline"})])
        self.assertEqual(len(p.inline_results[0]["issues"][0]["anchors"]), 2)

    def test_attachment_is_evidence_body_is_edit(self):
        a = [anchor("正文约定", "path:body/p1", "根据附件2第3条修改正文响应要求。", defect_type="inconsistency"),
             anchor("附件约定", "path:body/p2", "", role="evidence", defect_type="inconsistency")]
        p = build_review_presentation([result(opinion(a))])
        self.assertEqual([x["quoted_text"] for x in p.inline_results[0]["issues"][0]["anchors"]], ["正文约定"])

    def test_same_missing_field_deduplicates_across_criteria(self):
        a = result()
        b = result(opinion([anchor(comment="请补全此处实际地址。")]), criterion_id="C8")
        p = build_review_presentation([a, b])
        self.assertEqual(len(p.inline_results), 1)
        self.assertEqual(len(p.inline_results[0]["issues"][0]["source_issue_keys"]), 2)

    def test_different_defects_same_paragraph_are_not_deduplicated(self):
        a = result(opinion([anchor(defect_type="wording", comment="删除重复字词。")]), criterion="语病")
        b = result(opinion([anchor(defect_type="other", comment="补充责任主体。")]), criterion="语病", criterion_id="C9")
        self.assertEqual(len(build_review_presentation([a, b]).inline_results), 2)

    def test_missing_field_cannot_pass_without_removal_reason(self):
        old = opinion([anchor(), anchor("电话：#电话#", "path:body/p2", "电话未填写。")])
        ledger = FindingLedger(old)
        new = opinion()
        self.assertFalse(ledger.review(new, {"status": "PASS", "removals": []}, 1))
        missing = ledger.missing(new)
        self.assertEqual(len(missing), 1)
        self.assertTrue(ledger.review(new, {"status": "PASS", "removals": [
            {"finding_id": missing[0]["finding_id"], "reason": "该发现被证明不属于当前标准。"}]}, 2))

    def test_unknown_removal_and_retained_retraction_cannot_pass(self):
        o = opinion()
        ledger = FindingLedger(o)
        self.assertFalse(ledger.review(o, {"status": "PASS", "removals": [{"finding_id": "bad", "reason": "x"}]}, 1))
        key = ledger.current(o)[0]["finding_id"]
        self.assertFalse(ledger.review(o, {"status": "PASS", "removals": [{"finding_id": key, "reason": "原判断错误"}]}, 2))

    def test_pending_findings_and_lost_fields_remain_visible(self):
        ledger = FindingLedger(opinion([anchor(), anchor("邮箱：#邮箱#", "path:body/p2", "邮箱未填写。")]))
        r = result(quality_status="pending_review", unresolved_issues=ledger.unresolved_issues(opinion()))
        p = build_review_presentation([r])
        self.assertEqual(p.inline_results, [])
        self.assertEqual(len(p.global_issues), 2)
        self.assertTrue(all("待复核" in text for text in p.global_issues))

    def test_prompt_contracts(self):
        prompts = load_prompts(PROMPTS_PATH)
        self.assertNotIn("merged_by_party", prompts["PLANNER_STRATEGY_PROMPT"])
        for name in ["SUB_AGENT_BASE_PROMPT", "REFLECTOR_SYSTEM_PROMPT"]:
            for term in ["附件", "正文", "书名号", "edit", "evidence", "字段"]:
                self.assertIn(term, prompts[name])
        self.assertIn("missing_previous_findings", prompts["REFLECTOR_SYSTEM_PROMPT"])
        self.assertIn("pending_review", prompts["SUMMARIZER_SYSTEM_PROMPT"])

    def test_invalid_planner_scan_combination_rejected_before_retrieval(self):
        with self.assertRaises(ValidationError):
            ExtractedRoute(steps=[{"method": "region", "regions": ["frontmatter"]}, {"method": "scan"}])
        ExtractedRoute(steps=[{"method": "rule", "terms": ["主体"]}, {"method": "scan"}])


class WriterTests(unittest.TestCase):
    def test_exact_scopes_existing_comments_and_evidence_only(self):
        with tempfile.TemporaryDirectory() as temp:
            src, first, out = [Path(temp) / name for name in ("source.docx", "first.docx", "out.docx")]
            doc = Document()
            doc.add_paragraph("地址：#地址#")
            doc.add_paragraph("电话：#电话#")
            doc.add_paragraph("附件约定")
            doc.save(src)
            old = result(opinion([anchor(comment="原有人工意见", defect_type="other")]))
            DocxReportGenerator._generate_docx_with_comments(str(src), [old], str(first))
            new = result(opinion([anchor(), anchor("电话：#电话#", "path:body/p2", "电话未填写。"),
                                  anchor("附件约定", "path:body/p3", "", role="evidence")]))
            DocxReportGenerator._generate_docx_with_comments(str(first), [new], str(out))
            ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
            with zipfile.ZipFile(out) as z:
                comments = ET.fromstring(z.read("word/comments.xml"))
                texts = ["".join(c.itertext()) for c in comments]
                self.assertEqual(len(comments), 3)
                self.assertTrue(any("原有人工意见" in text for text in texts))
                body = ET.fromstring(z.read("word/document.xml"))
                paragraphs = body.findall(".//w:body/w:p", ns)
                self.assertEqual(len(paragraphs[0].findall(".//w:commentRangeStart", ns)), 2)
                self.assertEqual(len(paragraphs[1].findall(".//w:commentRangeStart", ns)), 1)
                self.assertEqual(len(paragraphs[2].findall(".//w:commentRangeStart", ns)), 0)


class ReflectionGateTests(unittest.IsolatedAsyncioTestCase):
    async def test_max_rejection_is_pending_and_last_round_is_not_unreviewed_rewrite(self):
        orchestrator = OrchestratorAgent()
        orchestrator._retrieve_context = AsyncMock(return_value=("context", 0))
        orchestrator._get_tools = AsyncMock(return_value=[])
        orchestrator._build_missing_text_review_notes = AsyncMock(return_value="")
        orchestrator._review_with_fallback = AsyncMock(return_value={"status": "REJECT", "feedback": "错误", "removals": []})
        with patch("agents.orchestrator.MAX_REFLECTION_ROUNDS", 1), \
             patch("agents.orchestrator.chat_until_valid_json", new=AsyncMock(return_value=(opinion(), "{}"))) as chat, \
             patch.object(orchestrator, "_validate_and_repair_anchors", return_value={}):
            r = await orchestrator.execute_single_criterion({"id": "C3", "criterion": "字段完整性"})
        self.assertEqual(r["quality_status"], "pending_review")
        self.assertEqual(chat.await_count, 1)
        self.assertEqual(len(r["quality_history"]), 1)


if __name__ == "__main__":
    unittest.main()
