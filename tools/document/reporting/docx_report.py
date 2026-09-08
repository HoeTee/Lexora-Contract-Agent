import json
import os
import re
import sys
import copy
import shutil
from contextlib import suppress
from dataclasses import dataclass
from datetime import datetime, timedelta, timezone
from pathlib import Path
from docx import Document
from docx.document import Document as DocxDocument
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.opc.packuri import PackURI
from docx.table import Table, _Cell
from docx.text.paragraph import Paragraph
from docx.text.run import Run
from lxml import etree

from config import DOCX_COMMENT_INCLUDE_CRITERION
from tools.document.docx_anchor_index import paragraph_anchor_id, table_anchor_id


SUMMARY_COMMENT_AUTHOR = "AI 审查总结"
REVIEW_COMMENT_AUTHOR = "AI 条款审查"
COMMENTS_RELTYPE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/comments"
RISK_LEVEL_LABELS = {
    "high": "高",
    "medium": "中",
    "low": "低",
}


@dataclass(frozen=True)
class ParagraphAnchor:
    """Resolved contract paragraph that can host a Word comment."""

    paragraph: Paragraph
    path: str
    normalized_text: str


@dataclass(frozen=True)
class CharRef:
    """A character in review text mapped back to its original DOCX run."""

    run: Run
    offset: int


@dataclass(frozen=True)
class ParagraphTextView:
    """Paragraph text view built with accepted-revision semantics."""

    paragraph: Paragraph
    path: str
    text: str
    char_map: list[CharRef]
    xml_anchor_id: str = ""


@dataclass(frozen=True)
class RunRange:
    """Selected character range inside one original DOCX run."""

    run: Run
    start_offset: int
    end_offset: int


@dataclass(frozen=True)
class TextMatch:
    """Quoted text match inside a paragraph text view."""

    start_index: int
    end_index: int
    matched_text: str


@dataclass(frozen=True)
class TextAnchor:
    """Resolved original-DOCX anchor for an AI comment."""

    match: TextMatch
    paragraph: Paragraph
    path: str
    run_ranges: list[RunRange]


def deduplicate_review_results_by_range(
    contract_path: str,
    results: list[dict],
) -> tuple[list[dict], dict[str, int]]:
    """Drop cross-source duplicate edit anchors by resolved DOCX character range.

    Results from the mandatory profile win only when an uploaded criterion resolves
    to the exact same paragraph path, start index, and end index. Unresolved anchors,
    summary-only findings, and different character ranges are preserved.
    """
    doc = Document(contract_path)
    candidates: list[tuple[int, int, int, tuple[str, int, int], bool]] = []
    for result_index, result in enumerate(results):
        mandatory = result.get("review_source") == "mandatory"
        for issue_index, issue in enumerate(result.get("issues") or []):
            for anchor_index, anchor in enumerate(issue.get("anchors") or []):
                if anchor.get("role", "edit") != "edit":
                    continue
                resolved = DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                    doc,
                    str(anchor.get("xml_anchor_type") or ""),
                    str(anchor.get("xml_anchor_id") or ""),
                    str(anchor.get("quoted_text") or ""),
                )
                if resolved is None:
                    continue
                key = (
                    resolved.path,
                    resolved.match.start_index,
                    resolved.match.end_index,
                )
                candidates.append((result_index, issue_index, anchor_index, key, mandatory))

    grouped: dict[tuple[str, int, int], list[tuple[int, int, int, bool]]] = {}
    suppressed: set[tuple[int, int, int]] = set()
    for result_index, issue_index, anchor_index, key, mandatory in candidates:
        grouped.setdefault(key, []).append((result_index, issue_index, anchor_index, mandatory))
    for group in grouped.values():
        if not any(item[3] for item in group):
            continue
        suppressed.update(item[:3] for item in group if not item[3])

    deduplicated = copy.deepcopy(results)
    removed_issues = 0
    removed_anchors = 0
    for result_index, result in enumerate(deduplicated):
        kept_issues = []
        for issue_index, issue in enumerate(result.get("issues") or []):
            original_anchors = issue.get("anchors") or []
            had_edit = any(anchor.get("role", "edit") == "edit" for anchor in original_anchors)
            kept_anchors = [
                anchor
                for anchor_index, anchor in enumerate(original_anchors)
                if (result_index, issue_index, anchor_index) not in suppressed
            ]
            removed_anchors += len(original_anchors) - len(kept_anchors)
            has_edit = any(anchor.get("role", "edit") == "edit" for anchor in kept_anchors)
            if had_edit and not has_edit:
                removed_issues += 1
                continue
            issue["anchors"] = kept_anchors
            kept_issues.append(issue)
        result["issues"] = kept_issues
    return deduplicated, {
        "resolved_edit_anchors": len(candidates),
        "removed_duplicate_anchors": removed_anchors,
        "removed_duplicate_issues": removed_issues,
    }


class DocxReportGenerator:
    """Generates DOCX reports, including annotated copies of original contracts."""

    # ==================== DOCX Comment Helpers ====================

    @staticmethod
    def _normalize_text(text: str) -> str:
        """Normalize text for fuzzy matching: strip whitespace, punctuation variations."""
        if not text:
            return ""
        text = re.sub(r'^>\s*', '', text, flags=re.MULTILINE)
        text = re.sub(r'[\s\u3000]+', '', text)
        text = text.replace('\uff08', '(').replace('\uff09', ')').replace('\uff1a', ':')
        text = text.replace('\uff0c', ',').replace('\u3002', '.').replace('\uff1b', ';')
        return text

    @staticmethod
    def _iter_blocks(parent):
        """Yield Paragraph/Table blocks in document order for a document or cell."""
        for _, block in DocxReportGenerator._iter_indexed_blocks(parent):
            yield block

    @staticmethod
    def _iter_indexed_blocks(parent):
        """Yield real XML child indexes with commentable blocks.

        Anchor paths are created from the position in ``w:body``/``w:tc``.  The
        position must not be recomputed after unsupported XML children (for
        example content controls) have been filtered out.
        """
        if isinstance(parent, DocxDocument):
            parent_element = parent.element.body
        elif isinstance(parent, _Cell):
            parent_element = parent._tc
        else:
            raise TypeError(f"Unsupported block parent: {type(parent)!r}")

        for child_index, child in enumerate(parent_element.iterchildren(), start=1):
            if isinstance(child, CT_P):
                yield child_index, Paragraph(child, parent)
            elif isinstance(child, CT_Tbl):
                yield child_index, Table(child, parent)

    @staticmethod
    def _collect_paragraph_anchors(parent, path_prefix: str = "body/") -> list[ParagraphAnchor]:
        """Collect commentable paragraphs from body text and nested table cells."""
        anchors: list[ParagraphAnchor] = []
        seen_elements: set[int] = set()

        for block_index, block in DocxReportGenerator._iter_indexed_blocks(parent):
            if isinstance(block, Paragraph):
                block_text = block.text.strip()
                normalized_text = DocxReportGenerator._normalize_text(block_text)
                element_id = id(block._element)
                if not normalized_text or element_id in seen_elements:
                    continue
                seen_elements.add(element_id)
                anchors.append(
                    ParagraphAnchor(
                        paragraph=block,
                        path=f"{path_prefix}p{block_index}",
                        normalized_text=normalized_text,
                    )
                )
                continue

            table_path = f"{path_prefix}tbl{block_index}/"
            for row_index, row in enumerate(block.rows, start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    anchors.extend(
                        DocxReportGenerator._collect_paragraph_anchors(
                            cell,
                            path_prefix=f"{table_path}r{row_index}c{cell_index}/",
                        )
                    )

        return anchors

    @staticmethod
    def _xml_local_name(element) -> str:
        """Return an XML element local name without its namespace."""
        return etree.QName(element).localname

    @staticmethod
    def _has_ancestor(element, local_names: set[str]) -> bool:
        """Return whether an XML element is nested under one of local_names."""
        parent = element.getparent()
        while parent is not None:
            if DocxReportGenerator._xml_local_name(parent) in local_names:
                return True
            parent = parent.getparent()
        return False

    @staticmethod
    def _build_paragraph_text_view(
        paragraph: Paragraph,
        path: str,
        xml_anchor_id: str = "",
    ) -> ParagraphTextView | None:
        """Build review-visible text and map each character to the original run."""
        chars: list[str] = []
        char_map: list[CharRef] = []

        for run_element in paragraph._element.iter(qn("w:r")):
            if DocxReportGenerator._has_ancestor(run_element, {"del", "moveFrom"}):
                continue

            run = Run(run_element, paragraph)
            run_offset = 0
            for text_element in run_element.iter():
                local_name = DocxReportGenerator._xml_local_name(text_element)
                if local_name == "delText":
                    continue
                if local_name != "t":
                    continue
                if DocxReportGenerator._has_ancestor(text_element, {"del", "moveFrom"}):
                    continue

                text = text_element.text or ""
                for char in text:
                    chars.append(char)
                    char_map.append(CharRef(run=run, offset=run_offset))
                    run_offset += 1

        if not chars:
            return None
        return ParagraphTextView(
            paragraph=paragraph,
            path=path,
            text="".join(chars),
            char_map=char_map,
            xml_anchor_id=xml_anchor_id,
        )

    @staticmethod
    def _collect_paragraph_text_views(parent, path_prefix: str = "body/") -> list[ParagraphTextView]:
        """Collect paragraph text views from body text and nested table cells."""
        views: list[ParagraphTextView] = []
        seen_elements: set[int] = set()

        for block_index, block in DocxReportGenerator._iter_indexed_blocks(parent):
            if isinstance(block, Paragraph):
                element_id = id(block._element)
                if element_id in seen_elements:
                    continue
                seen_elements.add(element_id)
                view = DocxReportGenerator._build_paragraph_text_view(
                    block,
                    f"{path_prefix}p{block_index}",
                    paragraph_anchor_id(block, f"{path_prefix}p{block_index}"),
                )
                if view and DocxReportGenerator._normalize_text(view.text):
                    views.append(view)
                continue

            table_path = f"{path_prefix}tbl{block_index}/"
            for row_index, row in enumerate(block.rows, start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    views.extend(
                        DocxReportGenerator._collect_paragraph_text_views(
                            cell,
                            path_prefix=f"{table_path}r{row_index}c{cell_index}/",
                        )
                    )

        return views

    @staticmethod
    def _collect_table_text_views(table: Table, table_path: str) -> list[ParagraphTextView]:
        views: list[ParagraphTextView] = []
        for row_index, row in enumerate(table.rows, start=1):
            for cell_index, cell in enumerate(row.cells, start=1):
                views.extend(
                    DocxReportGenerator._collect_paragraph_text_views(
                        cell,
                        path_prefix=f"{table_path}/r{row_index}c{cell_index}/",
                    )
                )
        return views

    @staticmethod
    def _collect_text_views_for_xml_anchor(
        parent,
        anchor_type: str,
        anchor_id: str,
        path_prefix: str = "body/",
    ) -> list[ParagraphTextView]:
        if not anchor_type or not anchor_id:
            return []

        for block_index, block in DocxReportGenerator._iter_indexed_blocks(parent):
            if isinstance(block, Paragraph):
                path = f"{path_prefix}p{block_index}"
                current_anchor_id = paragraph_anchor_id(block, path)
                if anchor_type == "paragraph" and current_anchor_id == anchor_id:
                    view = DocxReportGenerator._build_paragraph_text_view(
                        block,
                        path,
                        current_anchor_id,
                    )
                    return [view] if view else []
                continue

            table_path = f"{path_prefix}tbl{block_index}"
            current_table_anchor_id = table_anchor_id(block, table_path)
            if anchor_type == "table" and current_table_anchor_id == anchor_id:
                return DocxReportGenerator._collect_table_text_views(block, table_path)

            for row_index, row in enumerate(block.rows, start=1):
                for cell_index, cell in enumerate(row.cells, start=1):
                    nested_views = DocxReportGenerator._collect_text_views_for_xml_anchor(
                        cell,
                        anchor_type,
                        anchor_id,
                        path_prefix=f"{table_path}/r{row_index}c{cell_index}/",
                    )
                    if nested_views:
                        return nested_views

        return []

    @staticmethod
    def _run_ranges_from_char_refs(char_refs: list[CharRef]) -> list[RunRange]:
        """Collapse selected character refs into contiguous ranges per run."""
        if not char_refs:
            return []

        ranges: list[RunRange] = []
        current_run = char_refs[0].run
        start_offset = char_refs[0].offset
        previous_offset = char_refs[0].offset

        for ref in char_refs[1:]:
            if ref.run._r is current_run._r and ref.offset == previous_offset + 1:
                previous_offset = ref.offset
                continue

            ranges.append(
                RunRange(
                    run=current_run,
                    start_offset=start_offset,
                    end_offset=previous_offset + 1,
                )
            )
            current_run = ref.run
            start_offset = ref.offset
            previous_offset = ref.offset

        ranges.append(
            RunRange(
                run=current_run,
                start_offset=start_offset,
                end_offset=previous_offset + 1,
            )
        )
        return ranges

    @staticmethod
    def _text_anchor_from_indexes(
        view: ParagraphTextView,
        start_index: int,
        end_index: int,
    ) -> TextAnchor | None:
        if start_index < 0 or end_index <= start_index:
            return None
        selected_refs = view.char_map[start_index:end_index]
        run_ranges = DocxReportGenerator._run_ranges_from_char_refs(selected_refs)
        if not run_ranges:
            return None
        return TextAnchor(
            match=TextMatch(
                start_index=start_index,
                end_index=end_index,
                matched_text=view.text[start_index:end_index],
            ),
            paragraph=view.paragraph,
            path=view.path,
            run_ranges=run_ranges,
        )

    @staticmethod
    def _text_anchor_from_find(view: ParagraphTextView, needle: str) -> TextAnchor | None:
        if not needle:
            return None
        start = view.text.find(needle)
        if start < 0:
            return None
        return DocxReportGenerator._text_anchor_from_indexes(
            view,
            start,
            start + len(needle),
        )

    @staticmethod
    def _find_exact_text_range_anchor(
        views: list[ParagraphTextView],
        reference_text: str,
    ) -> TextAnchor | None:
        quoted = (reference_text or "").strip()
        if not quoted:
            return None
        for view in views:
            text_anchor = DocxReportGenerator._text_anchor_from_find(view, quoted)
            if text_anchor:
                return text_anchor
        normalized_quoted = DocxReportGenerator._normalize_text(quoted)
        if not normalized_quoted:
            return None
        for view in views:
            normalized_chars: list[str] = []
            original_indexes: list[int] = []
            for original_index, char in enumerate(view.text):
                normalized_char = DocxReportGenerator._normalize_text(char)
                if not normalized_char:
                    continue
                normalized_chars.append(normalized_char)
                original_indexes.extend([original_index] * len(normalized_char))
            normalized_text = "".join(normalized_chars)
            normalized_start = normalized_text.find(normalized_quoted)
            if normalized_start < 0:
                continue
            normalized_end = normalized_start + len(normalized_quoted) - 1
            if normalized_end >= len(original_indexes):
                continue
            text_anchor = DocxReportGenerator._text_anchor_from_indexes(
                view,
                original_indexes[normalized_start],
                original_indexes[normalized_end] + 1,
            )
            if text_anchor:
                return text_anchor
        return None

    @staticmethod
    def _find_text_range_anchor_in_xml_anchor(
        doc: Document,
        anchor_type: str,
        anchor_id: str,
        reference_text: str,
    ) -> TextAnchor | None:
        """Find the first quoted-text occurrence inside one DOCX XML anchor scope."""
        if not reference_text or not reference_text.strip():
            return None
        views = DocxReportGenerator._collect_text_views_for_xml_anchor(
            doc,
            anchor_type,
            anchor_id,
        )
        return DocxReportGenerator._find_exact_text_range_anchor(views, reference_text)

    @staticmethod
    def _get_document_start_anchor(doc: Document) -> ParagraphAnchor | None:
        """Return the first non-empty body paragraph that can host fallback comments."""
        anchors = DocxReportGenerator._collect_paragraph_anchors(doc)
        return anchors[0] if anchors else None

    @staticmethod
    def _get_document_end_anchor(doc: Document) -> ParagraphAnchor | None:
        """Return the last non-empty body paragraph for end-of-document issues."""
        anchors = DocxReportGenerator._collect_paragraph_anchors(doc)
        return anchors[-1] if anchors else None

    @staticmethod
    def _select_missing_issue_anchor(
        doc: Document,
        criterion: str,
        issue: dict,
        region_boundaries: dict[str, dict[str, str]] | None = None,
    ) -> tuple[ParagraphAnchor | None, str]:
        """Choose a deterministic placement for issues caused by missing text."""
        placement_text = " ".join(
            str(value or "")
            for value in (
                criterion,
                issue.get("check_point"),
                issue.get("issue_comment"),
            )
        )
        end_hints = ("正文最后", "正文结尾", "文末", "签署栏", "签字盖章", "以下无正文")
        if any(hint in placement_text for hint in end_hints):
            boundary = (region_boundaries or {}).get("body_content_end")
            if boundary:
                matched = DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                    doc,
                    boundary.get("xml_anchor_type", "paragraph"),
                    boundary.get("xml_anchor_id", ""),
                    boundary.get("quoted_text", ""),
                )
                if matched:
                    return matched, "body_content_end_missing_issue"
            return None, "body_content_end_unavailable"
        return DocxReportGenerator._get_document_start_anchor(doc), "document_start_missing_issue"

    @staticmethod
    def _paragraph_content_start_index(para_element) -> int:
        """Return the first valid child index after paragraph properties."""
        if len(para_element) and para_element[0].tag == qn('w:pPr'):
            return 1
        return 0

    @staticmethod
    def _split_run_at(run: Run, offset: int) -> Run | None:
        """Split a run in-place at offset, preserving style on the right run."""
        text = run.text
        if offset <= 0 or offset >= len(text):
            return None

        left_text = text[:offset]
        right_text = text[offset:]
        right_run_element = copy.deepcopy(run._r)
        run.text = left_text
        right_run = Run(right_run_element, run._parent)
        right_run.text = right_text
        run._r.addnext(right_run_element)
        return right_run

    @staticmethod
    def _split_run_range(run_range: RunRange) -> Run:
        """Split one run range so the returned run contains only selected text."""
        run = run_range.run
        if run_range.end_offset < len(run.text):
            DocxReportGenerator._split_run_at(run, run_range.end_offset)
        if run_range.start_offset > 0:
            selected_run = DocxReportGenerator._split_run_at(run, run_range.start_offset)
            if selected_run is not None:
                return selected_run
        return run

    @staticmethod
    def _set_run_highlight(run: Run, color: str = "yellow") -> None:
        """Apply Word highlight to a run without changing its existing text."""
        run_properties = run._r.get_or_add_rPr()
        highlight = run_properties.find(qn("w:highlight"))
        if highlight is None:
            highlight = OxmlElement("w:highlight")
            run_properties.append(highlight)
        highlight.set(qn("w:val"), color)

    @staticmethod
    def _insert_before(reference_element, new_element) -> None:
        parent = reference_element.getparent()
        parent.insert(parent.index(reference_element), new_element)

    @staticmethod
    def _insert_after(reference_element, new_element) -> None:
        parent = reference_element.getparent()
        parent.insert(parent.index(reference_element) + 1, new_element)

    @staticmethod
    def _add_comment_markers_to_text_range(anchor: TextAnchor, comment_id: int) -> None:
        selected_runs: list[Run] = []
        for run_range in anchor.run_ranges:
            selected_run = DocxReportGenerator._split_run_range(run_range)
            DocxReportGenerator._set_run_highlight(selected_run)
            selected_runs.append(selected_run)

        if not selected_runs:
            raise ValueError("No runs found for comment range")

        first_run_element = selected_runs[0]._r
        last_run_element = selected_runs[-1]._r

        range_start = OxmlElement('w:commentRangeStart')
        range_start.set(qn('w:id'), str(comment_id))
        DocxReportGenerator._insert_before(first_run_element, range_start)

        range_end = OxmlElement('w:commentRangeEnd')
        range_end.set(qn('w:id'), str(comment_id))
        DocxReportGenerator._insert_after(last_run_element, range_end)

        ref_run = OxmlElement('w:r')
        ref_rpr = OxmlElement('w:rPr')
        ref_style = OxmlElement('w:rStyle')
        ref_style.set(qn('w:val'), 'CommentReference')
        ref_rpr.append(ref_style)
        ref_run.append(ref_rpr)
        ref_mark = OxmlElement('w:commentReference')
        ref_mark.set(qn('w:id'), str(comment_id))
        ref_run.append(ref_mark)
        DocxReportGenerator._insert_after(range_end, ref_run)

    @staticmethod
    def _add_comment_markers_to_paragraph(anchor: ParagraphAnchor, comment_id: int) -> None:
        para_element = anchor.paragraph._element

        range_start = OxmlElement('w:commentRangeStart')
        range_start.set(qn('w:id'), str(comment_id))
        para_element.insert(
            DocxReportGenerator._paragraph_content_start_index(para_element),
            range_start,
        )

        range_end = OxmlElement('w:commentRangeEnd')
        range_end.set(qn('w:id'), str(comment_id))
        para_element.append(range_end)

        ref_run = OxmlElement('w:r')
        ref_rpr = OxmlElement('w:rPr')
        ref_style = OxmlElement('w:rStyle')
        ref_style.set(qn('w:val'), 'CommentReference')
        ref_rpr.append(ref_style)
        ref_run.append(ref_rpr)
        ref_mark = OxmlElement('w:commentReference')
        ref_mark.set(qn('w:id'), str(comment_id))
        ref_run.append(ref_mark)
        para_element.append(ref_run)

    @staticmethod
    def _build_issue_comment_text(
        comment_text: str,
        risk_level: str | None,
        criterion: str | None = None,
    ) -> str:
        """Build the Word comment body for a successfully anchored issue."""
        text = (comment_text or "").strip()
        risk_label = RISK_LEVEL_LABELS.get(str(risk_level or "").strip().lower())
        lines = []
        if risk_label:
            lines.append(f"风险等级：{risk_label}")
        if text:
            lines.append(text)

        criterion_text = (criterion or "").strip()
        if DOCX_COMMENT_INCLUDE_CRITERION and criterion_text:
            lines.extend(["", "审查要点：", criterion_text])

        return "\n".join(lines)

    @staticmethod
    def _build_summary_comment_text(summary_sections: dict | None) -> str:
        """Build the document-start summary comment from summarizer output."""
        if not summary_sections:
            return ""

        lines: list[str] = []
        overall_comment = str(summary_sections.get("overall_comment", "")).strip()
        priority_comments = summary_sections.get("priority_comments", [])
        opening_notices = summary_sections.get("opening_notices", [])

        if overall_comment:
            lines.append("总体审查结论：")
            lines.append(overall_comment)

        valid_priority_comments = [
            item.strip()
            for item in priority_comments
            if isinstance(item, str) and item.strip()
        ]
        if valid_priority_comments:
            if lines:
                lines.append("")
            lines.append("优先修改建议：")
            for index, item in enumerate(valid_priority_comments, start=1):
                lines.append(f"{index}. {item}")

        valid_opening_notices = [
            item.strip()
            for item in opening_notices
            if isinstance(item, str) and item.strip()
        ]
        if valid_opening_notices:
            if lines:
                lines.append("")
            lines.append("合同首部需补充的提示：")
            for index, item in enumerate(valid_opening_notices, start=1):
                lines.append(f"{index}. {item}")

        return "\n".join(lines)

    @staticmethod
    def _add_comments_to_doc(doc: Document, comments_data: list) -> Document:
        """Add Word comments to a document using OPC XML manipulation."""
        comments_part = None
        for rel in doc.part.rels.values():
            if rel.reltype == COMMENTS_RELTYPE:
                comments_part = rel.target_part
                break

        if comments_part is None:
            from docx.opc.part import Part

            comments_xml = (
                '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>'
                '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'
                ' xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">'
                '</w:comments>'
            )
            comments_part = Part(
                # python-docx expects OPC part names as PackURI objects; a raw
                # string breaks package traversal during doc.save().
                partname=PackURI("/word/comments.xml"),
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
                blob=comments_xml.encode('utf-8'),
                package=doc.part.package,
            )
            doc.part.relate_to(comments_part, COMMENTS_RELTYPE)

        comments_element = etree.fromstring(comments_part.blob)
        replaced_summary_ids: set[str] = set()
        for existing_comment in list(comments_element.findall(f".//{qn('w:comment')}")):
            if existing_comment.get(qn("w:author")) != SUMMARY_COMMENT_AUTHOR:
                continue
            raw_id = existing_comment.get(qn("w:id"))
            if raw_id is not None:
                replaced_summary_ids.add(raw_id)
            parent = existing_comment.getparent()
            if parent is not None:
                parent.remove(existing_comment)
        if replaced_summary_ids:
            marker_tags = {
                qn("w:commentRangeStart"),
                qn("w:commentRangeEnd"),
                qn("w:commentReference"),
            }
            for element in list(doc.element.iter()):
                if element.tag not in marker_tags or element.get(qn("w:id")) not in replaced_summary_ids:
                    continue
                parent = element.getparent()
                if parent is not None:
                    parent.remove(element)
        existing_ids: list[int] = []
        for existing_comment in comments_element.findall(f".//{qn('w:comment')}"):
            raw_id = existing_comment.get(qn("w:id"))
            try:
                existing_ids.append(int(raw_id))
            except (TypeError, ValueError):
                continue
        comment_id = max(existing_ids, default=-1) + 1
        beijing_tz = timezone(timedelta(hours=8))
        now_str = datetime.now(beijing_tz).isoformat(timespec="seconds")

        for cdata in comments_data:
            anchor = cdata.get('anchor')
            comment_text = cdata.get('comment_text', '')
            author = cdata.get('author', 'AI审查助手')

            if anchor is None or not comment_text:
                continue

            comment_el = etree.SubElement(comments_element, qn('w:comment'))
            comment_el.set(qn('w:id'), str(comment_id))
            comment_el.set(qn('w:author'), author)
            comment_el.set(qn('w:date'), now_str)

            comment_lines = comment_text.split('\n')
            for line in comment_lines:
                line = line.strip()
                p_el = etree.SubElement(comment_el, qn('w:p'))
                if not line:
                    continue
                r_el = etree.SubElement(p_el, qn('w:r'))
                t_el = etree.SubElement(r_el, qn('w:t'))
                t_el.set(qn('xml:space'), 'preserve')
                t_el.text = line

            if isinstance(anchor, TextAnchor):
                DocxReportGenerator._add_comment_markers_to_text_range(anchor, comment_id)
            else:
                DocxReportGenerator._add_comment_markers_to_paragraph(anchor, comment_id)

            comment_id += 1

        comments_blob = etree.tostring(
            comments_element,
            xml_declaration=True,
            encoding='UTF-8',
            standalone=True,
        )
        comments_part._blob = comments_blob
        if hasattr(comments_part, "_element"):
            comments_part._element = comments_element

        return doc

    @staticmethod
    def _generate_standard_docx_report(content: str, output_path: str) -> str | None:
        """Generate the standard DOCX report from markdown content."""
        try:
            from htmldocx import HtmlToDocx

            html_body = DocxReportGenerator._markdown_to_html(content)
            full_html = f"<html><body>{html_body}</body></html>"
            doc = Document()
            parser = HtmlToDocx()
            parser.add_html_to_document(full_html, doc)
            doc.save(output_path)
            return output_path
        except Exception as e:
            print(f"  WARNING: DOCX generation failed: {e}", file=sys.stderr)
            return None

    @staticmethod
    def _can_generate_annotated_docx(contract_path: str | None, results: list | None) -> bool:
        """Return whether the original contract can be used for annotated DOCX export."""
        return bool(
            contract_path
            and results
            and os.path.exists(contract_path)
            and contract_path.lower().endswith(".docx")
        )

    @staticmethod
    def _generate_docx_with_comments(
        contract_path: str,
        results: list,
        output_path: str,
        summary_sections: dict | None = None,
        annotation_events_path: str | None = None,
        region_boundaries: dict[str, dict[str, str]] | None = None,
    ) -> str:
        """Generate the annotated DOCX report using the original contract as the base."""

        try:
            shutil.copy2(contract_path, output_path)
            doc = Document(output_path)

            comments_data = []
            unmatched_comments = []
            annotation_events: list[dict] = []
            annotation_stats = {
                "total_issues": 0,
                "xml_anchor_matched": 0,
                "missing_text_fallback": 0,
                "unmatched": 0,
            }

            for result in results:
                criterion = result.get('criterion', '')
                cid = result.get('criterion_id', '')
                placement = result.get('placement') or {}
                placement_target = placement.get('target', '')

                for issue in result.get('issues', []):
                    annotation_stats["total_issues"] += 1
                    risk_level = issue.get('risk_level', '')
                    issue_id = issue.get('issue_id', '')
                    anchors = issue.get('anchors') or []

                    if not anchors:
                        if placement_target == 'body_end':
                            boundary = (region_boundaries or {}).get("body_content_end") or {}
                            fallback_anchor = (
                                DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                                    doc,
                                    boundary.get("xml_anchor_type", "paragraph"),
                                    boundary.get("xml_anchor_id", ""),
                                    boundary.get("quoted_text", ""),
                                )
                                if boundary
                                else None
                            )
                            fallback_strategy = (
                                "body_content_end_missing_issue" if fallback_anchor else "body_content_end_unavailable"
                            )
                        else:
                            fallback_anchor, fallback_strategy = (
                                DocxReportGenerator._select_missing_issue_anchor(
                                    doc,
                                    criterion,
                                    issue,
                                    region_boundaries,
                                )
                            )
                        if fallback_anchor:
                            comments_data.append({
                                'anchor': fallback_anchor,
                                'comment_text': DocxReportGenerator._build_issue_comment_text(
                                    issue.get('issue_comment', ''),
                                    risk_level,
                                    criterion=issue.get('criterion') or criterion,
                                ),
                                'author': REVIEW_COMMENT_AUTHOR,
                            })
                        else:
                            unmatched_comments.append({
                                'criterion_id': cid,
                                'criterion': criterion,
                                'issue_id': issue_id,
                                'risk_level': risk_level,
                                'comment_text': issue.get('issue_comment', ''),
                                'reference_text': '',
                                'author': REVIEW_COMMENT_AUTHOR,
                            })
                        annotation_stats["missing_text_fallback"] += 1
                        annotation_events.append({
                            "event": "annotation_missing_text",
                            "criterion_id": cid,
                            "issue_id": issue_id,
                            "placement": placement_target or None,
                            "status": "anchored" if fallback_anchor else "unmatched",
                            "reason": "anchors_empty",
                            "match_strategy": fallback_strategy,
                            "paragraph_path": fallback_anchor.path if fallback_anchor else None,
                            "risk_level": risk_level,
                            "comment_text": issue.get('issue_comment', ''),
                        })
                        continue

                    for anchor_item in anchors:
                        if anchor_item.get('role', 'edit') != 'edit':
                            continue
                        anchor_type = anchor_item.get('xml_anchor_type', '')
                        anchor_id = anchor_item.get('xml_anchor_id', '')
                        reference = anchor_item.get('quoted_text', '')
                        comment_text = anchor_item.get('comment_text', '')
                        matched_anchor = DocxReportGenerator._find_text_range_anchor_in_xml_anchor(
                            doc,
                            anchor_type,
                            anchor_id,
                            reference,
                        )

                        if matched_anchor:
                            comments_data.append({
                                'anchor': matched_anchor,
                                'comment_text': DocxReportGenerator._build_issue_comment_text(
                                    comment_text,
                                    risk_level,
                                    criterion=issue.get('criterion') or criterion,
                                ),
                                'author': REVIEW_COMMENT_AUTHOR,
                            })
                            annotation_stats["xml_anchor_matched"] += 1
                            annotation_events.append({
                                "event": "annotation_anchor_resolved",
                                "criterion_id": cid,
                                "issue_id": issue_id,
                                "status": "anchored",
                                "match_strategy": "xml_anchor_scoped_first_quoted_text",
                                "xml_anchor_type": anchor_type,
                                "xml_anchor_id": anchor_id,
                                "paragraph_path": matched_anchor.path,
                                "start_char": matched_anchor.match.start_index,
                                "end_char": matched_anchor.match.end_index,
                                "quoted_text": reference,
                                "matched_text": matched_anchor.match.matched_text,
                                "risk_level": risk_level,
                                "comment_text": comment_text,
                            })
                        else:
                            unmatched_comments.append({
                                'criterion_id': cid,
                                'criterion': criterion,
                                'issue_id': issue_id,
                                'risk_level': risk_level,
                                'comment_text': comment_text,
                                'reference_text': reference,
                                'author': REVIEW_COMMENT_AUTHOR,
                            })
                            annotation_stats["unmatched"] += 1
                            annotation_events.append({
                                "event": "annotation_anchor_unmatched",
                                "criterion_id": cid,
                                "issue_id": issue_id,
                                "status": "unmatched",
                                "reason": "xml_anchor_scoped_quoted_text_match_failed",
                                "xml_anchor_type": anchor_type,
                                "xml_anchor_id": anchor_id,
                                "quoted_text": reference,
                                "risk_level": risk_level,
                                "comment_text": comment_text,
                            })

            fallback_anchor = DocxReportGenerator._get_document_start_anchor(doc)
            summary_comment_text = DocxReportGenerator._build_summary_comment_text(summary_sections)
            fallback_count = 0
            if fallback_anchor and summary_comment_text:
                comments_data.insert(0, {
                    'anchor': fallback_anchor,
                    'comment_text': summary_comment_text,
                    'author': SUMMARY_COMMENT_AUTHOR,
                })
                fallback_count = 1

            skipped_count = len(unmatched_comments)

            DocxReportGenerator._add_comments_to_doc(doc, comments_data)

            doc.save(output_path)
            if annotation_events_path:
                annotation_path = Path(annotation_events_path)
                annotation_path.parent.mkdir(parents=True, exist_ok=True)
                annotation_path.write_text(
                    json.dumps(
                        {
                            "contract_path": contract_path,
                            "output_path": output_path,
                            "stats": annotation_stats,
                            "summary_comment_count": fallback_count,
                            "skipped_count": skipped_count,
                            "events": annotation_events,
                        },
                        ensure_ascii=False,
                        indent=2,
                    ),
                    encoding="utf-8",
                )
            print(
                f"  DOCX: {annotation_stats['xml_anchor_matched']} XML-anchored comments, "
                f"{fallback_count} document-start fallback comments, "
                f"{skipped_count} skipped",
                file=sys.stderr,
            )
            return output_path

        except Exception as e:
            print(f"  WARNING: DOCX comment generation failed: {e}", file=sys.stderr)
            import traceback
            traceback.print_exc(file=sys.stderr)
            if os.path.exists(output_path):
                with suppress(OSError):
                    os.unlink(output_path)
            return None

    # ==================== Markdown HTML Helper ====================

    @staticmethod
    def _markdown_to_html(md_content: str) -> str:
        """Convert markdown text to an HTML body fragment."""
        import markdown
        extensions = [
            'tables',
            'fenced_code',
            'codehilite',
            'toc',
            'nl2br',
            'sane_lists',
        ]
        extension_configs = {
            'codehilite': {'css_class': 'highlight', 'guess_lang': False},
        }
        html_body = markdown.markdown(
            md_content,
            extensions=extensions,
            extension_configs=extension_configs,
        )
        return html_body

    @staticmethod
    def generate_docx_report(
        content: str,
        contract_path: str = None,
        results: list = None,
        output_path: str = None,
    ) -> str | None:
        """Generate an annotated DOCX when possible, otherwise a standard DOCX report."""
        if DocxReportGenerator._can_generate_annotated_docx(contract_path, results):
            docx_path = DocxReportGenerator._generate_docx_with_comments(
                contract_path, results, output_path
            )
            if docx_path:
                return docx_path
            print(
                "  DOCX: annotated export failed, falling back to standard report DOCX.",
                file=sys.stderr,
            )

        return DocxReportGenerator._generate_standard_docx_report(content, output_path)

    @staticmethod
    def generate_annotated_docx(
        contract_path: str,
        results: list,
        output_path: str,
        summary_sections: dict | None = None,
        annotation_events_path: str | None = None,
        region_boundaries: dict[str, dict[str, str]] | None = None,
    ) -> str | None:
        """Generate only the annotated original-contract DOCX."""
        if not DocxReportGenerator._can_generate_annotated_docx(contract_path, results):
            print(
                "  DOCX: annotated export requires a DOCX contract and non-empty results.",
                file=sys.stderr,
            )
            return None
        return DocxReportGenerator._generate_docx_with_comments(
            contract_path,
            results,
            output_path,
            summary_sections=summary_sections,
            annotation_events_path=annotation_events_path,
            region_boundaries=region_boundaries,
        )


