"""Build the Word reading edition from the same source as the offline webpage."""

import argparse
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def clean(text):
    return text.replace("`", "").replace("**", "")


def font(style, size, name="Microsoft YaHei"):
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)


def table(document, rows):
    count = len(rows[0])
    widths = [2.0, 5.0] if count == 2 else [1.35, 2.8, 2.85]
    if rows[0][1] == "接口":
        widths = [1.1, 3.05, 2.85]
    result = document.add_table(rows=0, cols=count)
    result.autofit = False
    for column, width in zip(result.columns, widths):
        column.width = Inches(width)
    properties = result._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = OxmlElement("w:" + side)
        for key, value in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
            border.set(qn("w:" + key), value)
        borders.append(border)
    properties.append(borders)
    for index, values in enumerate(rows):
        row = result.add_row()
        row_props = row._tr.get_or_add_trPr()
        row_props.append(OxmlElement("w:cantSplit"))
        if index == 0:
            row_props.append(OxmlElement("w:tblHeader"))
        for cell, text, width in zip(row.cells, values, widths):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            props = cell._tc.get_or_add_tcPr()
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "17365D" if index == 0 else ("F3F5F7" if index % 2 else "FFFFFF"))
            props.append(shade)
            margins = OxmlElement("w:tcMar")
            for side, size in (("top", 95), ("bottom", 95), ("left", 130), ("right", 130)):
                margin = OxmlElement("w:" + side)
                margin.set(qn("w:w"), str(size))
                margin.set(qn("w:type"), "dxa")
                margins.append(margin)
            props.append(margins)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            run = paragraph.add_run(clean(text))
            run.font.size = Pt(10.5)
            if index == 0:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def diagram(document, kind, lines):
    # Word uses equivalent compact records instead of screenshotting web diagrams.
    if kind.startswith("timeline"):
        rows = [["阶段", "目标", "具体步骤"]]
        rows.extend([[clean(part.strip()) for part in line.split("|", 2)] for line in lines if line.strip()])
    else:
        rows = [["处理过程", "触发条件"]]
        for line in lines:
            if not line.strip():
                continue
            edge, separator, condition = line.partition(":")
            rows.append([edge.replace("->", "→").strip(), condition.strip() if separator else "按顺序处理"])
    table(document, rows)


def build(source, output):
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.7)
    section.left_margin = section.right_margin = Inches(0.75)
    normal = document.styles["Normal"]
    font(normal, 11)
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 24), ("Subtitle", 12), ("Heading 1", 18), ("Heading 2", 14)):
        font(document.styles[name], size)
    # The default template's Title style can carry a blue bottom border.
    for border in document.styles.element.xpath(".//w:pBdr"):
        border.getparent().remove(border)
    for name in ("Heading 1", "Heading 2"):
        document.styles[name].paragraph_format.keep_with_next = True
        document.styles[name].paragraph_format.space_before = Pt(12)
        document.styles[name].paragraph_format.space_after = Pt(7)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.add_paragraph("人体 Debug 互动系统设计", "Title")
    document.add_paragraph("产品流程与技术实现 阅读修订版", "Subtitle")
    lines = source.read_text(encoding="utf-8").splitlines()
    if lines[0] == "---":
        lines = lines[lines.index("---", 1) + 1:]
    index = 0
    prototype_mode = False
    first_prototype = True
    first_user_figure = True
    while index < len(lines):
        line = lines[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("## "):
            title = line[3:]
            if title.startswith("G "):
                landscape = document.add_section(WD_SECTION_START.NEW_PAGE)
                landscape.orientation = WD_ORIENT.LANDSCAPE
                landscape.page_width, landscape.page_height = Inches(11), Inches(8.5)
                landscape.top_margin = landscape.bottom_margin = Inches(0.5)
                landscape.left_margin = landscape.right_margin = Inches(0.6)
                prototype_mode = True
            paragraph = document.add_paragraph(title, "Heading 1")
            if not title.startswith(("A ", "G ")):
                paragraph.paragraph_format.page_break_before = True
            first_prototype = True
        elif line.startswith("### "):
            paragraph = document.add_paragraph(line[4:], "Heading 2")
            if prototype_mode:
                paragraph.paragraph_format.page_break_before = not first_prototype
                first_prototype = False
        elif line.startswith("```"):
            kind = line[3:]
            block = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                block.append(lines[index])
                index += 1
            if kind.startswith(("timeline", "flow")):
                diagram(document, kind, block)
            else:
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.line_spacing = 1.1
                run = paragraph.add_run("\n".join(block))
                run.font.name = "Consolas"
                run.font.size = Pt(10)
        elif line.startswith("|"):
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [part.strip() for part in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", part) for part in values):
                    rows.append(values)
                index += 1
            table(document, rows)
            continue
        elif line.startswith("!["):
            match = re.fullmatch(r"!\[([^]]+)\]\(([^)]+)\)", line)
            if not match:
                raise ValueError(f"Invalid figure: {line}")
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            # 1440x900 original images stay intact, one full figure per landscape page.
            height = 5.0 if first_user_figure else 5.35
            shape = paragraph.add_run().add_picture(str(source.parent / match[2]), height=Inches(height))
            first_user_figure = False
            shape._inline.docPr.set("descr", match[1])
        else:
            paragraph = document.add_paragraph(clean(line[2:] if line.startswith("- ") else line), "List Bullet" if line.startswith("- ") else "Normal")
            if prototype_mode:
                paragraph.paragraph_format.keep_with_next = True
        index += 1
    document.core_properties.title = "人体 Debug 互动系统设计"
    document.core_properties.subject = "产品流程与技术实现 阅读修订版"
    document.core_properties.author = ""
    document.save(output)
    print(f"Word reading edition: {output.resolve()}")
    print(f"Figures: {len(document.inline_shapes)}")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    build(args.source.resolve(), args.output.resolve())


if __name__ == "__main__":
    main()
