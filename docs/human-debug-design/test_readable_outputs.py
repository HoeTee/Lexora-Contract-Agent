"""Regression checks for the readable Word and single-file HTML deliverables."""

import base64
import re
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import html

from build_offline import validate_offline


ROOT = Path(__file__).resolve().parent
WORD = ROOT / "人体Debug互动系统设计文档 易读版.docx"


class ReadableOutputTests(unittest.TestCase):
    def test_webpage_has_every_original_figure_expanded(self):
        page = html.fromstring((ROOT / "index.html").read_text(encoding="utf-8"))
        validate_offline(page)
        self.assertFalse(page.xpath("//details"))
        figures = page.xpath("//img")
        self.assertEqual(len(figures), 25)
        for index, figure in enumerate(figures, 1):
            self.assertEqual(
                base64.b64decode(figure.get("src").split(",", 1)[1]),
                (ROOT / f"assets/prototype-{index:02}.png").read_bytes(),
            )
        self.assertNotIn("reference.html", page.text_content())

    def test_word_keeps_every_original_figure_and_heading(self):
        document = Document(WORD)
        self.assertEqual(len(document.inline_shapes), 25)
        self.assertEqual(sum(p.style.name == "Heading 1" for p in document.paragraphs), 8)
        self.assertEqual(
            sum(bool(re.match(r"^[UAC]\d{2} ", p.text)) for p in document.paragraphs), 25
        )
        self.assertFalse(document.styles.element.xpath(".//w:pBdr"))
        with ZipFile(WORD) as archive:
            media = {archive.read(name) for name in archive.namelist() if name.startswith("word/media/")}
        originals = {(ROOT / f"assets/prototype-{i:02}.png").read_bytes() for i in range(1, 26)}
        self.assertEqual(media, originals)

    def test_word_preserves_source_prose_and_table_cells(self):
        document = Document(WORD)
        text = "\n".join(p.text for p in document.paragraphs)
        text += "\n" + "\n".join(c.text for t in document.tables for r in t.rows for c in r.cells)
        source = (ROOT / "overview.md").read_text(encoding="utf-8").splitlines()
        source = source[source.index("---", 1) + 1:]
        in_code = False
        for line in source:
            if line.startswith("```"):
                in_code = not in_code
                continue
            if in_code or not line.strip() or line.startswith("!["):
                continue
            if line.startswith("|"):
                parts = [part.strip() for part in line.strip("|").split("|")]
                if all(re.fullmatch(r":?-+:?", part) for part in parts):
                    continue
            else:
                parts = [re.sub(r"^(?:#{2,3} |[-] )", "", line)]
            for part in parts:
                self.assertIn(part.replace("`", "").replace("**", ""), text)

    def test_offline_validator_rejects_missing_local_targets(self):
        for markup in (
            '<a href="reference.html">Old reference</a>',
            '<a href="#missing">Broken section</a>',
            '<img src="assets/prototype-01.png">',
            '<script src="https://example.com/library.js"></script>',
            '<p id="same"></p><p id="same"></p>',
        ):
            with self.subTest(markup=markup), self.assertRaises(ValueError):
                validate_offline(html.fromstring(markup))


if __name__ == "__main__":
    unittest.main()
